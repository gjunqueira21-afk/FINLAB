"""Testes do FinLab.

Rodar: python -m pytest finlab/tests -q

Cobrem o que quebra silenciosamente: extração contábil da CVM, escalas
(units, LPA em milhares), consistência dos múltiplos, curvas do score e o
proxy de LLM nos três formatos de API. O motor de valuation em JavaScript
tem teste próprio em finlab/tests/test_engine.py (roda no navegador).
"""

from __future__ import annotations

import json
import sys
import threading
import http.server
import socketserver
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from finlab.backend import agents, cvm, metrics, scoring, universe, valuation  # noqa: E402
from finlab.backend import market  # noqa: E402


# ---------------------------------------------------------------------------
# Universo
# ---------------------------------------------------------------------------

def test_universo_tem_90_acoes_sem_duplicatas():
    assert len(universe.UNIVERSE) == 90
    tickers = [c.ticker for c in universe.UNIVERSE]
    assert len(set(tickers)) == 90


def test_todo_setor_declarado_tem_empresa():
    usados = {c.sector for c in universe.UNIVERSE}
    assert usados == set(universe.SECTORS)


def test_metricas_do_setor_sao_conhecidas():
    for key, meta in universe.SECTORS.items():
        assert 1 <= len(meta["metrics"]) <= 4, key
        for m in meta["metrics"]:
            assert m in universe.METRIC_LABELS, (key, m)
            assert m in universe.METRIC_FORMAT, (key, m)


def test_units_tem_razao_maior_que_um():
    for ticker, ratio in universe.UNIT_RATIO.items():
        assert universe.get(ticker) is not None, ticker
        assert ratio >= 2
        assert ticker.endswith("11"), "só units terminam em 11"


# ---------------------------------------------------------------------------
# Extração da CVM
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not cvm.available(), reason="parquets da CVM ausentes")
def test_series_anuais_da_petrobras():
    dados = cvm.annual_series("009512")
    assert dados["years"], "sem exercícios"
    assert not dados["financial"]
    serie = dados["series"]
    receita = [v for v in serie["receita"] if v is not None]
    assert receita and min(receita) > 1e11, "receita da Petrobras acima de R$ 100 bi"
    # EBITDA precisa ser coerente com EBIT + D&A no mesmo ano
    for i, ano in enumerate(dados["years"]):
        ebit, da, ebitda = serie["ebit"][i], serie["depreciacao"][i], serie["ebitda"][i]
        if None in (ebit, da, ebitda):
            continue
        assert abs(ebitda - (ebit + abs(da))) < 1.0, ano


@pytest.mark.skipif(not cvm.available(), reason="parquets da CVM ausentes")
def test_capex_sempre_negativo_e_fcl_bate_com_fco_menos_capex():
    for ticker in ("VALE3", "WEGE3", "EQTL3", "SUZB3", "MGLU3"):
        comp = universe.get(ticker)
        dados = cvm.annual_series(comp.cd_cvm)
        serie = dados["series"]
        for i in range(len(dados["years"])):
            capex, fco, fcl = serie["capex"][i], serie["fco"][i], serie["fcl"][i]
            if capex is not None:
                assert capex <= 0, (ticker, dados["years"][i])
            if None not in (capex, fco, fcl):
                assert abs(fcl - (fco - abs(capex))) < 1.0, (ticker, dados["years"][i])


@pytest.mark.skipif(not cvm.available(), reason="parquets da CVM ausentes")
def test_bancos_sao_detectados_e_nao_ganham_ebitda():
    for ticker in ("ITUB4", "BBAS3", "BBDC4"):
        comp = universe.get(ticker)
        dados = cvm.annual_series(comp.cd_cvm)
        assert dados["financial"], ticker
        assert all(v is None for v in dados["series"]["ebitda"]), ticker
        assert all(v is None for v in dados["series"]["divida_liquida"]), ticker


@pytest.mark.skipif(not cvm.available(), reason="parquets da CVM ausentes")
def test_acoes_implicitas_no_lpa_corrigem_a_escala():
    """A conta 3.99 vem inflada em 1.000× pelo pipeline; o fallback corrige."""
    # Raia Drogasil: ~1,65 bilhão de ações.
    acoes = cvm.shares_from_eps("005258")
    assert acoes is not None
    assert 1e9 < acoes < 3e9, acoes


@pytest.mark.skipif(not cvm.available(), reason="parquets da CVM ausentes")
def test_capital_social_bate_com_ordem_de_grandeza_conhecida():
    casos = {"33.000.167/0001-01": (10e9, 15e9),   # Petrobras
             "33.592.510/0001-54": (4e9, 5e9),     # Vale
             "60.872.504/0001-23": (9e9, 13e9)}    # Itaú Unibanco
    for cnpj, (lo, hi) in casos.items():
        acoes = cvm.shares_outstanding(cnpj)
        assert acoes is not None and lo < acoes < hi, (cnpj, acoes)


# ---------------------------------------------------------------------------
# Métricas e múltiplos
# ---------------------------------------------------------------------------

def test_last_valid_pega_o_ultimo_nao_nulo():
    assert metrics.last_valid([1, None, 3, None], [2020, 2021, 2022, 2023]) == (3.0, 2022)
    assert metrics.last_valid([None, None], [2020, 2021]) == (None, None)
    assert metrics.last_valid([], []) == (None, None)


def test_cagr_exige_pontas_positivas():
    anos = [2020, 2021, 2022, 2023]
    assert metrics.cagr([100, 110, 121, 133.1], anos, 3) == pytest.approx(0.1, abs=1e-4)
    assert metrics.cagr([-10, 5, 8, 12], anos, 3) is None       # base negativa
    assert metrics.cagr([100, 110], [2022, 2023], 3) is None    # série curta


def test_div_protege_denominador():
    assert metrics.div(10, 2) == 5
    assert metrics.div(10, 0) is None
    assert metrics.div(None, 2) is None


def test_positive_years():
    assert metrics.positive_years([1, 2, -1, 4, 5]) == pytest.approx(0.8)
    assert metrics.positive_years([1, 2]) is None  # amostra insuficiente


def test_unit_divide_o_valor_de_mercado():
    fund = {"base": {"lucro_liquido": 1e9, "patrimonio_liquido": 5e9, "divida_liquida": None,
                     "ebitda": None, "ebit": None, "receita": None, "fcl": None},
            "indicadores": {}, "financial": True, "cnpj": None, "cd_cvm": None}
    serie = [("2026-01-02", 30.0), ("2026-01-03", 30.0)]

    snap_unit = metrics.market_snapshot("SANB11", serie, None, fund)
    snap_normal = metrics.market_snapshot("BBDC4", serie, None, fund)
    assert snap_unit["unit_ratio"] == 2
    assert snap_normal["unit_ratio"] == 1
    if snap_unit["shares"] and snap_normal["shares"]:
        assert snap_unit["shares_quote"] == snap_unit["shares"] / 2


def test_multiplos_sao_consistentes_com_os_insumos():
    fund = {"base": {"lucro_liquido": 200.0, "patrimonio_liquido": 1000.0,
                     "divida_liquida": 300.0, "ebitda": 400.0, "ebit": 350.0,
                     "receita": 2000.0, "fcl": 150.0},
            "indicadores": {"roe": 0.2, "mg_ebitda": 0.2, "nd_ebitda": 0.75},
            "financial": False}
    snap = {"market_cap": 2000.0, "shares_quote": 100.0, "shares": 100.0}
    mult = metrics.multiples(fund, snap, None)
    assert mult["pl"] == pytest.approx(10.0)
    assert mult["pvp"] == pytest.approx(2.0)
    assert mult["ev"] == pytest.approx(2300.0)
    assert mult["ev_ebitda"] == pytest.approx(5.75)
    assert mult["lpa"] == pytest.approx(2.0)
    assert mult["vpa"] == pytest.approx(10.0)


def test_financeira_nao_recebe_enterprise_value():
    fund = {"base": {"lucro_liquido": 100.0, "patrimonio_liquido": 500.0,
                     "divida_liquida": None, "ebitda": None, "ebit": None,
                     "receita": 900.0, "fcl": None},
            "indicadores": {}, "financial": True}
    mult = metrics.multiples(fund, {"market_cap": 1000.0, "shares_quote": 10.0}, None)
    assert mult["ev"] is None
    assert mult["ev_ebitda"] is None


def _fund_itub():
    # ITUB4, exercício 2025 na CVM: lucro 45,85 bi / PL 215,08 bi = ROE 21,3%.
    return {"base": {"lucro_liquido": 45.85e9, "patrimonio_liquido": 215.08e9,
                     "divida_liquida": None, "ebitda": None, "ebit": None,
                     "receita": 300e9, "fcl": None},
            "indicadores": {"roe": 45.85 / 215.08}, "financial": True, "last_year": 2025}


def test_multiplos_vem_da_brapi_quando_ela_tem():
    brapi = {"priceEarnings": 9.5, "dividendYield": 7.0,
             "defaultKeyStatistics": {"priceToBook": 2.2, "bookValue": 20.0,
                                      "trailingEps": 4.6},
             "financialData": {"returnOnEquity": 0.228}}
    snap = {"market_cap": 400e9, "shares_quote": 10e9, "price": 44.0}
    mult = metrics.multiples(_fund_itub(), snap, brapi)
    assert mult["pl"] == pytest.approx(9.5)
    assert mult["pvp"] == pytest.approx(2.2)
    assert mult["roe"] == pytest.approx(0.228)
    assert mult["dy"] == pytest.approx(0.07)
    # LPA/VPA na mesma janela do P/L e do P/VP da tela: preço ÷ múltiplo.
    assert mult["lpa"] == pytest.approx(44.0 / 9.5)
    assert mult["vpa"] == pytest.approx(20.0)
    assert mult["fontes"]["pl"] == mult["fontes"]["roe"] == "BRAPI"
    # Banco: nada de EV/EBITDA, nem que a BRAPI mande um.
    assert mult["ev_ebitda"] is None and mult["ev"] is None


def test_roe_da_brapi_em_pontos_percentuais_vira_fracao():
    # Mesmo número em pontos (22,8): a escala é a que bate com LPA ÷ VPA.
    brapi = {"defaultKeyStatistics": {"bookValue": 20.0, "trailingEps": 4.6},
             "financialData": {"returnOnEquity": 22.8}}
    snap = {"market_cap": 400e9, "shares_quote": 10e9, "price": 44.0}
    assert metrics.multiples(_fund_itub(), snap, brapi)["roe"] == pytest.approx(0.228)
    # Sem LPA/VPA da BRAPI, a referência é a conta da CVM.
    brapi = {"financialData": {"returnOnEquity": 22.8}}
    assert metrics.multiples(_fund_itub(), snap, brapi)["roe"] == pytest.approx(0.228)
    # ROE baixo em pontos (1,2%) não pode virar 120%: com LPA ÷ VPA da própria
    # BRAPI (mesma janela) a escala sai certa mesmo com o anual da CVM em 21%.
    brapi = {"defaultKeyStatistics": {"bookValue": 20.0, "trailingEps": 0.24},
             "financialData": {"returnOnEquity": 1.2}}
    assert metrics.multiples(_fund_itub(), snap, brapi)["roe"] == pytest.approx(0.012)
    # Sem LPA/VPA, a referência seguinte é a CVM de 12 meses (mesma janela).
    ltm = {"fim": "2026-06-30", "rotulo": "LTM 2T26",
           "campos": {"lucro_liquido": 2.6e9, "patrimonio_liquido": 215e9}}
    brapi = {"financialData": {"returnOnEquity": 1.2}}
    assert metrics.multiples(_fund_itub(), snap, brapi, ltm)["roe"] == pytest.approx(0.012)


def test_multiplos_caem_para_a_cvm_campo_a_campo():
    # BRAPI só com P/L: o resto vem do exercício da CVM, e a fonte diz isso.
    snap = {"market_cap": 430e9, "shares_quote": 10e9, "price": 43.0}
    mult = metrics.multiples(_fund_itub(), snap, {"priceEarnings": 9.0,
                                                  "financialData": {"returnOnEquity": None}})
    assert mult["fontes"]["pl"] == "BRAPI"
    assert mult["pvp"] == pytest.approx(430 / 215.08)
    assert mult["roe"] == pytest.approx(0.2132, abs=1e-4)
    assert mult["fontes"]["pvp"] == mult["fontes"]["roe"] == "CVM"
    assert mult["ano_cvm"] == 2025
    # Sem BRAPI nenhuma, tudo CVM — o comportamento de antes.
    mult = metrics.multiples(_fund_itub(), snap, None)
    assert mult["pl"] == pytest.approx(430 / 45.85)
    assert set(v for v in mult["fontes"].values() if v) == {"CVM"}


def test_multiplos_ignoram_lixo_da_brapi():
    snap = {"market_cap": 430e9, "shares_quote": 10e9, "price": 43.0}
    brapi = {"priceEarnings": "NaN", "dividendYield": "abc",
             "defaultKeyStatistics": {"priceToBook": 0},
             "financialData": {"returnOnEquity": float("inf")}}
    mult = metrics.multiples(_fund_itub(), snap, brapi)
    assert mult["pl"] == pytest.approx(430 / 45.85)
    assert mult["pvp"] == pytest.approx(430 / 215.08)
    assert mult["dy"] is None
    assert mult["fontes"]["pl"] == mult["fontes"]["pvp"] == mult["fontes"]["roe"] == "CVM"


def test_nao_financeira_usa_ev_ebitda_e_margem_da_brapi():
    fund = {"base": {"lucro_liquido": 200.0, "patrimonio_liquido": 1000.0,
                     "divida_liquida": 300.0, "ebitda": 400.0, "ebit": 350.0,
                     "receita": 2000.0, "fcl": 150.0},
            "indicadores": {"roe": 0.2, "mg_ebitda": 0.2, "nd_ebitda": 0.75},
            "financial": False, "last_year": 2025}
    snap = {"market_cap": 2000.0, "shares_quote": 100.0, "price": 20.0}
    brapi = {"financialData": {"enterpriseToEbitda": 6.1, "ebitdaMargins": 0.22,
                               "enterpriseValue": 2500.0}}
    mult = metrics.multiples(fund, snap, brapi)
    assert mult["ev_ebitda"] == pytest.approx(6.1)
    assert mult["mg_ebitda"] == pytest.approx(0.22)
    assert mult["ev"] == pytest.approx(2500.0)
    # Alavancagem continua da CVM.
    assert mult["nd_ebitda"] == pytest.approx(0.75)
    assert mult["fontes"]["nd_ebitda"] == "CVM"


def _ltm_itub():
    # 12 meses até o 2T26 pelos ITRs: lucro 50 bi sobre PL de 220 bi = 22,7%.
    return {"fim": "2026-06-30", "rotulo": "LTM 2T26",
            "campos": {"lucro_liquido": 50e9, "patrimonio_liquido": 220e9,
                       "receita": 310e9}}


def test_cada_fonte_escolhida_e_respeitada():
    snap = {"market_cap": 440e9, "shares_quote": 10e9, "price": 44.0}
    brapi = {"priceEarnings": 9.5, "financialData": {"returnOnEquity": 0.215}}
    por = metrics.multiplos_por_fonte(_fund_itub(), snap, brapi, _ltm_itub())
    assert set(por) == set(metrics.FONTES)

    ex = por["cvm_exercicio"]
    assert ex["roe"] == pytest.approx(45.85 / 215.08)
    assert ex["pl"] == pytest.approx(440 / 45.85)
    assert ex["lpa"] == pytest.approx(4.585)
    assert set(v for v in ex["fontes"].values() if v) <= {"CVM", "BRAPI"}
    assert ex["fontes"]["roe"] == "CVM"

    m12 = por["cvm_12m"]
    assert m12["roe"] == pytest.approx(50 / 220)
    assert m12["pl"] == pytest.approx(440 / 50)
    assert m12["pvp"] == pytest.approx(2.0)
    assert m12["lpa"] == pytest.approx(5.0)
    assert m12["fontes"]["roe"] == "CVM12" and m12["ltm_rotulo"] == "LTM 2T26"

    b = por["brapi"]
    assert b["pl"] == pytest.approx(9.5) and b["roe"] == pytest.approx(0.215)
    # BRAPI escolhida e sem P/VP: fica vazio, não pega o da CVM escondido.
    assert b["pvp"] is None and b["fontes"]["pvp"] is None

    # Automático: BRAPI, depois CVM 12 meses, depois exercício — campo a campo.
    a = por["auto"]
    assert a["pl"] == pytest.approx(9.5) and a["fontes"]["pl"] == "BRAPI"
    assert a["pvp"] == pytest.approx(2.0) and a["fontes"]["pvp"] == "CVM12"
    # LPA e VPA na janela do P/L e do P/VP que estão na tela.
    assert a["lpa"] == pytest.approx(44.0 / 9.5) and a["fontes"]["lpa"] == "BRAPI"
    assert a["vpa"] == pytest.approx(22.0) and a["fontes"]["vpa"] == "CVM12"


def test_sem_itr_a_fonte_cvm_12m_fica_vazia_e_o_auto_cai_para_o_exercicio():
    snap = {"market_cap": 440e9, "shares_quote": 10e9, "price": 44.0}
    por = metrics.multiplos_por_fonte(_fund_itub(), snap, None, None)
    assert all(por["cvm_12m"][k] is None for k in ("pl", "pvp", "roe", "lpa"))
    assert por["auto"]["roe"] == pytest.approx(45.85 / 215.08)
    assert por["auto"]["fontes"]["roe"] == "CVM"


def test_divida_sobre_ebitda_vem_da_cvm_mesmo_com_brapi_escolhida():
    fund = {"base": {"lucro_liquido": 200.0, "patrimonio_liquido": 1000.0,
                     "divida_liquida": 300.0, "ebitda": 400.0, "ebit": 350.0,
                     "receita": 2000.0, "fcl": 150.0},
            "indicadores": {}, "financial": False, "last_year": 2025}
    ltm = {"fim": "2026-06-30", "rotulo": "LTM 2T26",
           "campos": {"lucro_liquido": 220.0, "patrimonio_liquido": 1100.0,
                      "divida_liquida": 250.0, "ebitda": 500.0, "receita": 2200.0}}
    snap = {"market_cap": 2000.0, "shares_quote": 100.0, "price": 20.0}
    por = metrics.multiplos_por_fonte(fund, snap, {"priceEarnings": 9.0}, ltm)
    assert por["brapi"]["nd_ebitda"] == pytest.approx(0.5)
    assert por["brapi"]["fontes"]["nd_ebitda"] == "CVM12"
    assert por["cvm_exercicio"]["nd_ebitda"] == pytest.approx(0.75)
    assert por["cvm_12m"]["ev_ebitda"] == pytest.approx(2250 / 500)
    assert por["cvm_12m"]["mg_ebitda"] == pytest.approx(500 / 2200)


def test_brapi_multiplos_cai_para_papel_a_papel_quando_o_plano_recusa_lote(monkeypatch):
    import requests
    from finlab.backend import cache as cache_mod
    monkeypatch.setattr(market, "BRAPI_TOKEN", "t")
    monkeypatch.setattr(cache_mod, "get", lambda *a, **k: None)
    monkeypatch.setattr(cache_mod, "set", lambda *a, **k: None)

    class Resp:
        status_code = 403

    def recusa(*a, **k):
        raise requests.HTTPError(response=Resp())
    monkeypatch.setattr(market._SESSION, "get", recusa)
    monkeypatch.setattr(market, "brapi_fundamentals",
                        lambda tk: {"symbol": tk, "priceEarnings": 8.0})
    out = market.brapi_multiplos(["itub4", "bbdc4"])
    assert set(out) == {"ITUB4", "BBDC4"}

    # Rede fora do ar: não sai disparando uma consulta por papel.
    chamadas = []
    monkeypatch.setattr(market._SESSION, "get",
                        lambda *a, **k: (_ for _ in ()).throw(requests.ConnectionError()))
    monkeypatch.setattr(market, "brapi_fundamentals", lambda tk: chamadas.append(tk))
    assert market.brapi_multiplos(["ITUB4"]) == {}
    assert chamadas == []


# ---------------------------------------------------------------------------
# Score
# ---------------------------------------------------------------------------

def test_curva_interpola_e_satura():
    ancoras = [(0.0, 0.0), (0.10, 50.0), (0.20, 100.0)]
    assert scoring.curve(0.05, ancoras) == pytest.approx(25.0)
    assert scoring.curve(0.15, ancoras) == pytest.approx(75.0)
    assert scoring.curve(-1.0, ancoras) == 0.0      # abaixo do piso
    assert scoring.curve(9.0, ancoras) == 100.0     # acima do teto
    assert scoring.curve(None, ancoras) is None


def test_curva_de_alavancagem_premia_menos_divida():
    assert scoring.curve(0.5, scoring.ND_EBITDA) > scoring.curve(3.5, scoring.ND_EBITDA)
    assert scoring.curve(-0.5, scoring.ND_EBITDA) > 95.0     # caixa líquido
    assert scoring.curve(-2.0, scoring.ND_EBITDA) == 100.0   # caixa líquido alto satura


def test_score_fica_entre_0_e_100_e_reporta_cobertura():
    otima = {"roe": 0.30, "roic": 0.25, "mg_liquida": 0.25, "nd_ebitda": -0.5,
             "nd_equity": -0.2, "mg_ebitda": 0.45, "cagr_receita_3a": 0.25,
             "cagr_ebitda_3a": 0.25, "cash_conversion": 1.0, "fcf_margin": 0.18,
             "consistencia_lucro": 1.0}
    pessima = {k: (-0.2 if "cagr" in k or "margin" in k or k.startswith("mg") or k in
                   ("roe", "roic", "cash_conversion") else 6.0) for k in otima}
    pessima["consistencia_lucro"] = 0.0

    boa = scoring.score(otima)
    ruim = scoring.score(pessima)
    assert 90 <= boa["total"] <= 100
    assert 0 <= ruim["total"] <= 20
    assert boa["cobertura"] == pytest.approx(1.0)
    assert not boa["parcial"]


def test_score_com_dado_faltando_nao_e_punido_mas_marca_cobertura():
    parcial = scoring.score({"roe": 0.20})
    assert parcial["total"] is not None
    assert parcial["cobertura"] < 0.3
    assert parcial["parcial"] is True


def test_score_sem_indicador_algum_nao_inventa_nota():
    assert scoring.score({})["total"] is None
    assert scoring.grade(None) == "—"
    assert scoring.band(None) == "none"


def test_perfil_financeiro_usa_outros_pilares():
    pilares = {p["key"] for p in scoring.score({"roe": 0.2}, financial=True)["pilares"]}
    assert "caixa" not in pilares          # conversão de caixa não se aplica a banco
    assert "rentabilidade" in pilares


# ---------------------------------------------------------------------------
# Premissas de valuation
# ---------------------------------------------------------------------------

def _fund_teste():
    return {
        "sector": "INDUSTRIA_TECH", "financial": False,
        "base": {"divida_bruta": 1000.0, "divida_liquida": 500.0, "caixa": 500.0},
        "indicadores": {"cagr_receita_3a": 0.08},
        "series": {"fcl": [100.0, 120.0, 140.0], "ebit": [200.0, 220.0, 240.0]},
    }


def test_premissas_respeitam_wacc_maior_que_perpetuidade():
    macro = {"selic": {"value": 14.15}, "ipca": {"value": 4.6}, "cdi": {"value": 14.15},
             "prefixado_10a": {"value": 14.8}, "ntnb_10a": {"value": 8.1}}
    prem = valuation.assumptions(_fund_teste(), {"market_cap": 5000.0, "shares_quote": 100.0,
                                                 "price": 40.0}, macro)
    assert prem["g_terminal"] < prem["wacc"] - 0.01
    assert prem["rf"] == pytest.approx(0.148)
    assert prem["rf_modo"] == "pre10"
    assert set(prem["rf_opcoes"]) == {"pre10", "ntnb", "selic"}
    assert prem["fcf_base"] == pytest.approx(120.0)     # média dos 3 últimos
    assert prem["ebit_normalizado"] == pytest.approx(220.0)
    assert len(prem["growth"]) == 5


def test_premissas_sem_curva_caem_para_a_selic():
    prem = valuation.assumptions(_fund_teste(), {"market_cap": 5000.0, "shares_quote": 100.0},
                                 {"selic": {"value": 14.15}, "ipca": {"value": 4.6}})
    assert prem["rf_modo"] == "selic"
    assert prem["rf"] == pytest.approx(0.1415)


def test_dcf_nao_se_aplica_a_instituicao_financeira():
    fund = _fund_teste()
    fund["financial"] = True
    prem = valuation.assumptions(fund, {"market_cap": 1.0, "shares_quote": 1.0}, {})
    assert prem["aplicavel"] is False
    assert "financeira" in prem["motivo_nao_aplicavel"].lower()


def test_estrutura_de_capital_vem_do_mercado_quando_ha_dado():
    prem = valuation.assumptions(_fund_teste(), {"market_cap": 3000.0, "shares_quote": 100.0}, {})
    assert prem["wd"] == pytest.approx(1000.0 / 4000.0)
    prem_sem = valuation.assumptions(_fund_teste(), {}, {})
    assert prem_sem["wd"] == pytest.approx(0.25)


# ---------------------------------------------------------------------------
# Utilidades de mercado
# ---------------------------------------------------------------------------

def test_numero_pt_br():
    assert market._pt_number("14,15%") == pytest.approx(14.15)
    assert market._pt_number("5.1177") == pytest.approx(5.1177)
    assert market._pt_number("175.335") == pytest.approx(175335.0)
    assert market._pt_number("1.234.567") == pytest.approx(1234567.0)
    assert market._pt_number("") is None
    assert market._pt_number(None) is None


def test_performance_calcula_janelas_e_respeita_tolerancia():
    from datetime import date, timedelta
    hoje = date(2026, 7, 27)
    serie = []
    for i in range(400, -1, -1):
        d = hoje - timedelta(days=i)
        serie.append((d.isoformat(), 100.0 * (1.0005 ** (400 - i))))
    perf = market.performance(serie, hoje)
    assert perf["price"] == pytest.approx(serie[-1][1])
    for janela in ("day", "week", "m3", "m12", "ytd"):
        assert perf[janela] is not None, janela
        assert perf[janela] > 0

    # Série curta: janelas longas não podem ser inventadas.
    curta = serie[-20:]
    perf_curta = market.performance(curta, hoje)
    assert perf_curta["week"] is not None
    assert perf_curta["m3"] is None
    assert perf_curta["m12"] is None
    assert perf_curta["ytd"] is None


def test_performance_sem_serie():
    assert market.performance([])["price"] is None


def test_performance_com_buracos_e_fim_de_semana():
    """A busca por data virou bisect sobre texto ISO; precisa achar o mesmo ponto.

    O caso que importa é a data-alvo que NÃO existe na série (fim de semana,
    feriado, pregão sem negócio): a janela tem de cair no último fechamento
    anterior ao alvo, nunca no seguinte.
    """
    from datetime import date

    serie = [("2026-01-02", 100.0), ("2026-01-05", 110.0), ("2026-01-06", 120.0),
             ("2026-04-06", 130.0), ("2026-07-06", 140.0)]
    # 2026-01-03 e 04 são fim de semana: a janela de 3 meses a partir de 06/04
    # tem de usar como base o fechamento de 05/01, não o de 06/01. O numerador
    # é sempre o último fechamento da série.
    perf = market.performance(serie, date(2026, 4, 6))
    assert perf["m3"] == pytest.approx(140.0 / 110.0 - 1)
    # Alvo anterior ao primeiro ponto: sem base, sem retorno inventado.
    assert market.performance(serie[-2:], date(2026, 7, 6))["m12"] is None


def test_pulse_recorta_o_historico_sem_perder_a_janela_do_painel(monkeypatch):
    """O blob de cotações era de 16 MB e o painel só consome ~2 anos.

    O corte tem de deixar de fora o que está além de ANOS_DE_HISTORICO e
    preservar tudo o que as telas leem (500 fechamentos e a janela de 12 meses).
    """
    from datetime import date, timedelta

    hoje = date.today()
    velha = (hoje - timedelta(days=365 * (market.ANOS_DE_HISTORICO + 2))).isoformat()
    recente = (hoje - timedelta(days=30)).isoformat()
    linhas = [
        {"label": "TESTE3", "data_referencia": velha, "preco_fechamento": "10"},
        {"label": "TESTE3", "data_referencia": recente, "preco_fechamento": "20"},
    ]
    monkeypatch.setattr(market, "_pulse_csv", lambda *a, **kw: linhas)
    monkeypatch.setattr(market.cache, "memoize", lambda key, ttl, producer: producer())

    serie = market.pulse_prices()["TESTE3"]
    assert [d for d, _ in serie] == [recente]


class _RespFake:
    def __init__(self, status, payload):
        self.status_code, self._payload, self.text = status, payload, str(payload)

    def json(self):
        return self._payload

    def close(self):
        pass


def test_xai_busca_usa_agent_tools_no_responses(monkeypatch):
    """O Live Search por search_parameters morreu (o provedor devolve 410).

    Com busca pedida, a chamada vai DIRETO ao /v1/responses com as ferramentas
    em `tools` — não passa pelo /chat/completions, onde a busca não existe.
    As fontes de `citations` entram no fim do texto.
    """
    chamadas = []

    def fake_post(url, **kw):
        chamadas.append((url, kw.get("json") or {}))
        return _RespFake(200, {"output_text": "contexto levantado",
                               "citations": ["https://x.com/post/1"]})

    monkeypatch.setattr(agents.requests, "post", fake_post)
    texto = agents.chat("xai", "k", "grok-4-fast", "sys", "usr", buscar=True)

    assert texto.startswith("contexto levantado")
    assert "https://x.com/post/1" in texto
    assert [u for u, _ in chamadas] == ["https://api.x.ai/v1/responses"]
    corpo = chamadas[0][1]
    assert corpo["tools"] == agents.FERRAMENTAS_BUSCA
    assert "input" in corpo and "messages" not in corpo
    assert "search_parameters" not in corpo


def test_xai_sem_busca_cai_para_o_responses_quando_recusado(monkeypatch):
    """Modelo novo pode não ser servido no /chat/completions: cai para o
    /v1/responses. Um 401 não repete — chave rejeitada não melhora tentando."""
    chamadas = []

    def fake_post(url, **kw):
        chamadas.append((url, kw.get("json") or {}))
        if url.endswith("/chat/completions"):
            return _RespFake(404, {"error": "model not found"})
        return _RespFake(200, {"output_text": "resposta"})

    monkeypatch.setattr(agents.requests, "post", fake_post)
    assert agents.chat("xai", "k", "grok-4.5", "sys", "usr") == "resposta"
    assert [u for u, _ in chamadas] == ["https://api.x.ai/v1/chat/completions",
                                        "https://api.x.ai/v1/responses"]
    # Sem busca pedida, nada de tools no corpo.
    assert "tools" not in chamadas[1][1]

    chamadas.clear()
    monkeypatch.setattr(agents.requests, "post",
                        lambda url, **kw: (chamadas.append(url), _RespFake(401, {}))[1])
    with pytest.raises(agents.LLMError):
        agents.chat("xai", "k", "grok-4.5", "sys", "usr")
    assert len(chamadas) == 1


def test_chat_conversa_com_busca_e_personas_novas(monkeypatch):
    """O chat também sabe buscar: o Agente de Contexto abre a rodada por lá.

    E as três vozes novas (contexto, cetico, moderador) têm persona própria —
    sem entrada em CHAT_PERSONAS, a fala sairia com o prompt genérico da mesa.
    """
    for chave in ("contexto", "cetico", "moderador"):
        assert chave in agents.CHAT_PERSONAS, chave

    chamadas = []

    def fake_post(url, **kw):
        chamadas.append((url, kw.get("json") or {}))
        return _RespFake(200, {"output_text": "radar do dia"})

    monkeypatch.setattr(agents.requests, "post", fake_post)
    texto = agents.chat_conversa("xai", "k", "grok-4-fast", "CTX", [],
                                 "o que estão falando?", "contexto", buscar=True)
    assert texto == "radar do dia"
    assert chamadas[0][0].endswith("/v1/responses")
    assert chamadas[0][1]["tools"] == agents.FERRAMENTAS_BUSCA


# ---------------------------------------------------------------------------
# Proxy de LLM
# ---------------------------------------------------------------------------

class _MockHandler(http.server.BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _responde(self, payload, codigo=200):
        data = json.dumps(payload).encode()
        self.send_response(codigo)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        self.server.recebido[self.path.split("?")[0]] = {"headers": dict(self.headers)}

        if self.path.startswith("/models/openai"):
            self._responde({"data": [{"id": "gpt-teste"}, {"id": "text-embedding-3"},
                                     {"id": "whisper-1"}, {"id": "gpt-teste"}]})
        elif self.path.startswith("/models/anthropic"):
            self._responde({"data": [{"id": "claude-teste", "display_name": "Claude Teste"}]})
        elif self.path.startswith("/models/google"):
            self._responde({"models": [
                {"name": "models/gemini-teste", "supportedGenerationMethods": ["generateContent"]},
                {"name": "models/embedding-001", "supportedGenerationMethods": ["embedContent"]},
            ]})
        elif self.path.startswith("/models/vazio"):
            self._responde({"data": []})
        elif self.path.startswith("/models/quebrado"):
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(b"nao sou json")
        elif self.path.startswith("/models/status/"):
            self.send_response(int(self.path.rsplit("/", 1)[-1]))
            self.end_headers()
            self.wfile.write(b"erro simulado")
        else:
            self._responde({"erro": "rota desconhecida"}, 404)

    def do_POST(self):
        tamanho = int(self.headers.get("Content-Length", 0))
        corpo = json.loads(self.rfile.read(tamanho) or b"{}")
        self.server.recebido[self.path] = {"body": corpo, "headers": dict(self.headers)}

        if self.path.startswith("/vazio/raciocinio"):
            # modelo de raciocínio: content vazio, texto no campo à parte
            payload = {"choices": [{"message": {"content": "",
                                                "reasoning": "VEIO-DO-RACIOCINIO"}}]}
        elif self.path.startswith("/vazio/truncado"):
            payload = {"choices": [{"message": {"content": ""},
                                    "finish_reason": "length"}]}
        elif self.path.startswith("/vazio/nulo"):
            payload = {"choices": [{"message": {"content": None}}]}
        elif self.path.startswith("/vazio/anthropic"):
            payload = {"content": [], "stop_reason": "max_tokens"}
        elif self.path.startswith("/vazio/gemini"):
            payload = {"candidates": [{"content": {"parts": []},
                                       "finishReason": "SAFETY"}]}
        elif self.path.startswith("/openai"):
            payload = {"choices": [{"message": {"content": "OK-OPENAI"}}]}
        elif self.path.startswith("/anthropic"):
            payload = {"content": [{"type": "text", "text": "OK-"},
                                   {"type": "text", "text": "ANTHROPIC"}]}
        elif self.path.startswith("/status/"):
            codigo = int(self.path.rsplit("/", 1)[-1])
            self.send_response(codigo)
            self.end_headers()
            self.wfile.write(b"erro simulado")
            return
        else:
            payload = {"candidates": [{"content": {"parts": [{"text": "OK-GEMINI"}]}}]}

        data = json.dumps(payload).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


@pytest.fixture()
def mock_llm(monkeypatch):
    servidor = socketserver.TCPServer(("127.0.0.1", 0), _MockHandler)
    servidor.recebido = {}
    threading.Thread(target=servidor.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{servidor.server_address[1]}"

    monkeypatch.setitem(agents.PROVIDERS["openrouter"], "url", base + "/openai")
    monkeypatch.setitem(agents.PROVIDERS["anthropic"], "url", base + "/anthropic")
    monkeypatch.setitem(agents.PROVIDERS["google"], "url", base + "/gemini")
    yield servidor, base
    servidor.shutdown()


def test_proxy_llm_nos_tres_formatos(mock_llm):
    servidor, _ = mock_llm
    assert agents.chat("openrouter", "sk-x", "m", "SYS", "USER") == "OK-OPENAI"
    assert agents.chat("anthropic", "sk-x", "m", "SYS", "USER") == "OK-ANTHROPIC"
    assert agents.chat("google", "sk-x", "m", "SYS", "USER") == "OK-GEMINI"

    aberto = servidor.recebido["/openai"]
    assert aberto["headers"]["Authorization"] == "Bearer sk-x"
    assert aberto["body"]["messages"][0]["content"] == "SYS"

    claude = servidor.recebido["/anthropic"]
    assert claude["headers"]["x-api-key"] == "sk-x"
    assert claude["body"]["system"] == "SYS"

    gemini_path = next(p for p in servidor.recebido if p.startswith("/gemini"))
    gemini = servidor.recebido[gemini_path]
    assert gemini["headers"]["x-goog-api-key"] == "sk-x"


def test_proxy_llm_traduz_erros(mock_llm, monkeypatch):
    _, base = mock_llm
    for codigo, trecho in ((401, "rejeitada"), (429, "Limite"), (500, "HTTP 500")):
        monkeypatch.setitem(agents.PROVIDERS["openrouter"], "url", f"{base}/status/{codigo}")
        with pytest.raises(agents.LLMError, match=trecho):
            agents.chat("openrouter", "sk-x", "m", "s", "u")


def test_proxy_llm_exige_chave_e_provedor_valido():
    with pytest.raises(agents.LLMError, match="Chave"):
        agents.chat("openrouter", "", "m", "s", "u")
    with pytest.raises(agents.LLMError, match="desconhecido"):
        agents.chat("nao-existe", "k", "m", "s", "u")


def test_contexto_da_tela_de_acoes_carrega_a_lista_inteira():
    overview = {
        "rows": [
            {"rank": 1, "ticker": "PRIO3", "sector": "OIL", "score": 87.6,
             "price": 40.0, "financial": False, "market_cap": 34e9,
             "perf": {"m12": 0.35, "ytd": 0.12},
             "multiples": {"pl": 9.9, "pvp": 2.3, "dy": 0.051, "roe": 0.22,
                           "ev_ebitda": 4.2, "nd_ebitda": 0.4},
             "valor": {"epv_por_acao": 38.4, "epv_upside": -0.04, "lpa": 4.05},
             "porte": {"receita": 25e9, "lucro_liquido": 7e9},
             "qualidade": {"mg_ebitda": 0.62, "mg_liquida": 0.28, "roic": 0.19,
                           "cagr_receita_3a": 0.31, "cagr_lucro_3a": 0.22}},
            {"rank": 2, "ticker": "ITUB4", "sector": "FIN", "score": 80.0,
             "price": None, "financial": True,
             "perf": {"m12": None, "ytd": None},
             "multiples": {"pl": 8.0, "pvp": 1.8, "dy": None, "roe": 0.21,
                           "nd_ebitda": None},
             "valor": {"epv_por_acao": None, "epv_upside": None, "lpa": 3.7}},
        ],
        "sector_stats": {"OIL": {"n": 7, "score": 68.0, "pl": 6.1, "pvp": 1.2,
                                 "dy": 0.08, "roe": 0.15}},
    }
    setores = {"OIL": {"label": "Petróleo"}, "FIN": {"label": "Bancos"}}
    macro = {"selic": {"value": "14,15%"}}

    ctx = agents.contexto_lista_acoes(overview, macro, setores)
    assert "TELA ABERTA" in ctx and "2 ações" in ctx
    # a linha carrega o dossiê resumido: valor, múltiplos, qualidade e porte
    assert ("1. PRIO3 (Petróleo) nota 87,6 | R$ 40,00 | EPV R$ 38,40 (dist -4,0%) | "
            "LPA 4,05 | 12m +35,0% · YTD +12,0% | P/L 9,9x · P/VP 2,3x · "
            "EV/EBITDA 4,2x · DY +5,1% · ROE +22,0% · ROIC +19,0% · DL/EBITDA 0,4x") in ctx
    assert "mg EBITDA +62,0% · mg líq +28,0% | CAGR rec +31,0% · lucro +22,0%" in ctx
    assert "receita R$ 25,00 bi · lucro R$ 7,00 bi · mkt cap R$ 34,00 bi" in ctx
    # banco: sem cotação não vira número; EPV não se aplica, mas o LPA sim
    assert "2. ITUB4 (Bancos) nota 80,0 | sem cotação | n/a (financeira) | LPA 3,70" in ctx

    # e a instrução de garimpo, que é o que autoriza a mesa a ordenar/filtrar
    # em vez de dizer que precisa abrir empresa por empresa
    assert "VOCÊ TEM O PAINEL INTEIRO AQUI" in ctx
    assert "não peça para abrir empresa por empresa" in ctx
    assert "\"Perto do EPV\" é |dist| pequena" in ctx
    assert "DL/EBITDA n/a" in ctx
    assert "Petróleo (7 ações): nota 68,0 · P/L 6,1x" in ctx
    assert "SELIC: 14,15%" in ctx


def test_contexto_da_tela_de_etfs_filtra_sem_negocio_e_corta_tese():
    payload = {"rows": [
        {"ticker": "BOVA11", "categoria": "INDICES_BR", "taxa_adm": 0.1,
         "liquidez": 1.8e9, "pl": 2.1e10, "price": 130.0,
         "perf": {"m12": 0.15, "ytd": 0.1}, "tese": "x" * 200},
        {"ticker": "MORTO11", "categoria": "INDICES_BR", "taxa_adm": None,
         "liquidez": 0, "pl": None, "price": None, "perf": {}, "tese": "sem negócio"},
    ]}
    cats = {"INDICES_BR": {"label": "Índices Brasil"}}
    ctx = agents.contexto_lista_etfs(payload, {}, cats)
    assert "2 ETFs" in ctx and "(1 sem negócios recentes ficaram fora" in ctx
    assert "BOVA11 (Índices Brasil) taxa 0,10% | liq R$ 1,8 bi | PL R$ 21,0 bi" in ctx
    assert "MORTO11" not in ctx
    assert "x" * 107 + "…" in ctx and "x" * 120 not in ctx   # tese truncada


def test_contexto_da_tela_de_bdrs_converte_setor_e_dy():
    payload = {"rows": [
        {"ticker": "AAPL34", "us_ticker": "AAPL", "name": "Apple",
         "sector": "TECHNOLOGY", "price": 61.5, "dy": 0.005, "liquidez": 1.2e7,
         "perf": {"m12": 0.3, "ytd": -0.02}},
    ]}
    setores = {"TECHNOLOGY": {"label": "Information Technology"}}
    ctx = agents.contexto_lista_bdrs(payload, {}, setores)
    assert "1 BDRs" in ctx
    assert "AAPL34 (AAPL) Apple · Information Technology | R$ 61,50 | " \
           "12m +30,0% | YTD -2,0% | DY +0,5% | liq R$ 12,0 mi" in ctx


def test_leitor_da_cvm_aceita_itr_e_degrada_sem_ele(tmp_path, monkeypatch):
    """O sufixo _dfp era fixo (achado 00.3): o pipeline gera *_itr.parquet e o
    painel não lia. Com um ITR sintético, latest_quarter devolve o trimestre
    mais recente; sem arquivo, devolve None sem quebrar nada."""
    import pandas as pd

    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        # sem arquivos: nada de ITR, nada de exceção
        assert cvm.quarterly_available() is False
        assert cvm.latest_quarter("009512") is None

        linhas = []
        for fim, receita, lucro in (("2026-03-31", 100.0, 10.0),
                                    ("2026-06-30", 220.0, 25.0)):
            linhas += [
                {"CD_CVM": "009512", "DENOM_CIA": "PETRO", "CNPJ_CIA": "x",
                 "DT_FIM_EXERC": pd.Timestamp(fim), "DT_INI_EXERC": pd.Timestamp("2026-01-01"),
                 "ANO_REFER": 2026, "CD_CONTA": "3.01",
                 "DS_CONTA": "Receita de Venda de Bens e/ou Serviços",
                 "VL_CONTA_AJUSTADO": receita},
                {"CD_CVM": "009512", "DENOM_CIA": "PETRO", "CNPJ_CIA": "x",
                 "DT_FIM_EXERC": pd.Timestamp(fim), "DT_INI_EXERC": pd.Timestamp("2026-01-01"),
                 "ANO_REFER": 2026, "CD_CONTA": "3.11",
                 "DS_CONTA": "Lucro/Prejuízo Consolidado do Período",
                 "VL_CONTA_AJUSTADO": lucro},
            ]
        # janela avulsa (2T isolado) NÃO pode ser confundida com o acumulado
        linhas.append({"CD_CVM": "009512", "DENOM_CIA": "PETRO", "CNPJ_CIA": "x",
                       "DT_FIM_EXERC": pd.Timestamp("2026-06-30"),
                       "DT_INI_EXERC": pd.Timestamp("2026-04-01"),
                       "ANO_REFER": 2026, "CD_CONTA": "3.01",
                       "DS_CONTA": "Receita de Venda de Bens e/ou Serviços",
                       "VL_CONTA_AJUSTADO": 120.0})
        pd.DataFrame(linhas).to_parquet(tmp_path / "dre_itr.parquet", index=False)
        cvm.limpar_cache()

        assert cvm.quarterly_available() is True
        q = cvm.latest_quarter("009512")
        assert q == {"fim": "2026-06-30", "receita": 220.0, "lucro": 25.0}
        # outra empresa segue sem dado, sem exceção
        assert cvm.latest_quarter("999999") is None
    finally:
        cvm.limpar_cache()


def test_downloader_revalida_quando_a_origem_muda(tmp_path, monkeypatch):
    """A CVM republica exercícios retroativamente; pular só porque o arquivo
    existe servia dado velho em silêncio (achado 00.4)."""
    import sys as _sys
    _sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "valuation_cvm"))
    dl = pytest.importorskip("src.cvm_downloader",
                             reason="dependências do pipeline (tqdm) ausentes")

    arq = tmp_path / "dfp.zip"
    arq.write_bytes(b"conteudo-antigo")

    class _Resp:
        def __init__(self, headers, status=200):
            self.headers = headers
            self.status_code = status

    # origem mais nova (Last-Modified no futuro) -> rebaixa
    monkeypatch.setattr(dl.requests, "head",
                        lambda *a, **k: _Resp({"Last-Modified": "Wed, 01 Jan 2225 00:00:00 GMT"}))
    assert dl._remote_is_newer("http://x/dfp.zip", arq) is True

    # origem antiga -> mantém o cache
    monkeypatch.setattr(dl.requests, "head",
                        lambda *a, **k: _Resp({"Last-Modified": "Wed, 01 Jan 2020 00:00:00 GMT"}))
    assert dl._remote_is_newer("http://x/dfp.zip", arq) is False

    # sem Last-Modified: decide pelo tamanho
    monkeypatch.setattr(dl.requests, "head",
                        lambda *a, **k: _Resp({"Content-Length": "999"}))
    assert dl._remote_is_newer("http://x/dfp.zip", arq) is True
    monkeypatch.setattr(dl.requests, "head",
                        lambda *a, **k: _Resp({"Content-Length": str(arq.stat().st_size)}))
    assert dl._remote_is_newer("http://x/dfp.zip", arq) is False

    # rede fora ou sem cabeçalho: fica com o local (comportamento antigo)
    def _boom(*a, **k):
        raise dl.requests.exceptions.ConnectionError("offline")
    monkeypatch.setattr(dl.requests, "head", _boom)
    assert dl._remote_is_newer("http://x/dfp.zip", arq) is False


def test_historico_corrompido_nao_derruba_o_painel(tmp_path, monkeypatch):
    """Duas instâncias do painel gravando junto corromperam o history.csv e
    TODA página de empresa passou a devolver 500. Uma linha ruim tem de ser
    pulada, não virar exceção."""
    arq = tmp_path / "history.csv"
    arq.write_text(
        "PETR4,2024-01-02,30.5\n"
        + "LIXO," + ("x" * 200000) + ",1\n"          # campo gigante
        + "VALE3,2024-01-02,nao-e-numero\n"          # preço inválido
        + "SO,DUAS,COLUNAS,DEMAIS\n"                 # colunas a mais
        + ",,\n"                                     # linha vazia
        + "VALE3,2024-01-03,61.25\n",
        encoding="utf-8")
    monkeypatch.setattr(market, "HISTORY_FILE", arq)

    hist = market._load_local_history()
    assert hist["PETR4"] == {"2024-01-02": 30.5}
    assert hist["VALE3"] == {"2024-01-03": 61.25}   # a linha ruim sumiu, a boa ficou
    assert "SO" not in hist and "" not in hist


def test_historico_e_gravado_de_forma_atomica(tmp_path, monkeypatch):
    arq = tmp_path / "history.csv"
    monkeypatch.setattr(market, "HISTORY_FILE", arq)
    market._save_local_history({"PETR4": {"2024-01-02": 30.5}})

    assert arq.read_text(encoding="utf-8").strip() == "PETR4,2024-01-02,30.500000"
    # nenhum temporário deixado para trás
    assert [p.name for p in tmp_path.iterdir()] == ["history.csv"]
    assert market._load_local_history() == {"PETR4": {"2024-01-02": 30.5}}


def test_status_das_fontes_diagnostica_o_token(monkeypatch):
    monkeypatch.setattr(market, "_probe", lambda k: {"brapi": True, "yahoo": True,
                                                     "pulseflat": False}[k])
    monkeypatch.setattr(market, "BRAPI_TOKEN", "abcd1234567890xyz")
    monkeypatch.setattr(market, "source_label", lambda: "BRAPI")

    st = market.provider_status()
    assert set(st) == {"brapi", "yahoo", "pulseflat"}
    assert st["brapi"]["configured"] and st["brapi"]["ok"] and st["brapi"]["em_uso"]
    assert st["yahoo"]["ok"] and not st["yahoo"]["precisa_token"]
    assert st["pulseflat"]["ok"] is False

    # o token é identificável mas nunca aparece inteiro
    mascara = st["brapi"]["token_mascarado"]
    assert "abcd1234567890xyz" not in mascara
    assert mascara.startswith("abcd") and "17 caracteres" in mascara

    # e o painel diz onde procurou o arquivo
    assert st["brapi"]["env_path"].endswith(".env")


def test_status_sem_token_nao_vaza_mascara(monkeypatch):
    monkeypatch.setattr(market, "_probe", lambda k: k != "brapi")
    monkeypatch.setattr(market, "BRAPI_TOKEN", "")
    monkeypatch.setattr(market, "source_label", lambda: "PulseFlat (B3/Yahoo D-1)")

    st = market.provider_status()
    assert st["brapi"]["configured"] is False
    assert st["brapi"]["token_mascarado"] == ""
    assert st["pulseflat"]["em_uso"] is True

    # token curto vira só bolinhas, sem revelar o tamanho útil
    monkeypatch.setattr(market, "BRAPI_TOKEN", "abc")
    assert market.provider_status()["brapi"]["token_mascarado"] == "•••"


def test_diagnostico_das_fontes_nao_entra_no_cache(monkeypatch):
    """O cache vive em disco e sobrevive ao restart. Se o estado do token
    fosse memoizado junto, quem acabasse de configurar o BRAPI_TOKEN veria
    'rodando sem token' pelo resto do TTL."""
    from finlab.backend import app as app_mod

    chamadas = {"n": 0}

    def status_falso():
        chamadas["n"] += 1
        return {"brapi": {"configured": chamadas["n"] > 1, "ok": True}}

    monkeypatch.setattr(app_mod.market, "provider_status", status_falso)
    monkeypatch.setattr(app_mod.market, "source_label", lambda: f"fonte-{chamadas['n']}")

    payload = {"rows": [1, 2, 3]}
    primeiro = app_mod._com_diagnostico(payload)
    segundo = app_mod._com_diagnostico(payload)

    assert primeiro["rows"] == [1, 2, 3] and segundo["rows"] == [1, 2, 3]
    # a segunda leitura enxerga o token que acabou de ser configurado
    assert primeiro["providers"]["brapi"]["configured"] is False
    assert segundo["providers"]["brapi"]["configured"] is True
    assert primeiro["source"] != segundo["source"]
    # e o payload original (o que fica no cache) segue limpo
    assert "providers" not in payload and "source" not in payload


def test_lista_de_modelos_le_a_api_do_provedor(mock_llm, monkeypatch):
    servidor, base = mock_llm
    monkeypatch.setitem(agents._MODEL_ENDPOINTS, "openrouter",
                        (base + "/models/openai", "bearer"))
    monkeypatch.setitem(agents._MODEL_ENDPOINTS, "anthropic",
                        (base + "/models/anthropic", "anthropic"))
    monkeypatch.setitem(agents._MODEL_ENDPOINTS, "google",
                        (base + "/models/google", "google"))

    r = agents.list_models("openrouter", "sk-x")
    assert r["fonte"] == "api" and r["aviso"] is None
    # embedding e whisper não servem para chat; duplicata some
    assert r["models"] == ["gpt-teste"]
    assert servidor.recebido["/models/openai"]["headers"]["Authorization"] == "Bearer sk-x"

    r = agents.list_models("anthropic", "sk-x")
    assert r["models"] == ["claude-teste"]
    assert servidor.recebido["/models/anthropic"]["headers"]["x-api-key"] == "sk-x"

    r = agents.list_models("google", "sk-x")
    # o prefixo "models/" some e o que só faz embedding fica de fora
    assert r["models"] == ["gemini-teste"]
    assert servidor.recebido["/models/google"]["headers"]["x-goog-api-key"] == "sk-x"


def test_lista_de_modelos_cai_no_catalogo_local_sem_quebrar(mock_llm, monkeypatch):
    _, base = mock_llm
    catalogo = agents.PROVIDERS["openrouter"]["models"]

    sem_chave = agents.list_models("openrouter", "")
    assert sem_chave["models"] == catalogo
    assert sem_chave["fonte"] == "catálogo local" and "Sem chave" in sem_chave["aviso"]

    casos = [("/models/status/401", "401"), ("/models/status/500", "HTTP 500"),
             ("/models/quebrado", "inesperada"), ("/models/vazio", "não listou")]
    for rota, trecho in casos:
        monkeypatch.setitem(agents._MODEL_ENDPOINTS, "openrouter", (base + rota, "bearer"))
        r = agents.list_models("openrouter", "sk-x")
        assert r["models"] == catalogo, rota
        assert trecho in r["aviso"], rota

    # provedor fora do ar: nada de exceção vazando para a tela
    monkeypatch.setitem(agents._MODEL_ENDPOINTS, "openrouter",
                        ("http://127.0.0.1:9/models", "bearer"))
    r = agents.list_models("openrouter", "sk-x")
    assert r["models"] == catalogo and "Não foi possível" in r["aviso"]

    with pytest.raises(agents.LLMError, match="desconhecido"):
        agents.list_models("nao-existe", "sk-x")


def test_conversa_leva_historico_e_contexto_nos_tres_formatos(mock_llm):
    servidor, _ = mock_llm
    hist = [{"role": "user", "content": "e a margem?"},
            {"role": "assistant", "content": "subiu"},
            {"role": "system", "content": "   "}]

    assert agents.chat_conversa(
        "openrouter", "sk-x", "m", "CTX", hist, "e a dívida?") == "OK-OPENAI"
    corpo = servidor.recebido["/openai"]["body"]
    assert corpo["messages"][0]["role"] == "system"
    assert "CTX" in corpo["messages"][0]["content"]
    # mensagem em branco não vira turno; a pergunta entra por último
    assert [m["role"] for m in corpo["messages"][1:]] == ["user", "assistant", "user"]
    assert corpo["messages"][-1]["content"] == "e a dívida?"

    assert agents.chat_conversa(
        "anthropic", "sk-x", "m", "CTX", hist, "e a dívida?") == "OK-ANTHROPIC"
    corpo = servidor.recebido["/anthropic"]["body"]
    assert "CTX" in corpo["system"]
    assert [m["role"] for m in corpo["messages"]] == ["user", "assistant", "user"]

    assert agents.chat_conversa(
        "google", "sk-x", "m", "CTX", hist, "e a dívida?") == "OK-GEMINI"
    rota = next(p for p in servidor.recebido if p.startswith("/gemini"))
    corpo = servidor.recebido[rota]["body"]
    assert "CTX" in corpo["systemInstruction"]["parts"][0]["text"]
    # o Gemini chama o papel do assistente de "model"
    assert [c["role"] for c in corpo["contents"]] == ["user", "model", "user"]


def test_conversa_nao_devolve_bolha_vazia(mock_llm, monkeypatch):
    """O bug do balão em branco: modelo de raciocínio ou resposta truncada
    devolviam string vazia, e a tela mostrava uma bolha sem nada dentro."""
    _, base = mock_llm

    # texto no campo de raciocínio é aproveitado em vez de virar vazio
    monkeypatch.setitem(agents.PROVIDERS["openrouter"], "url", base + "/vazio/raciocinio")
    assert agents.chat_conversa(
        "openrouter", "k", "m", "CTX", [], "e ai?") == "VEIO-DO-RACIOCINIO"

    # sem texto em lugar nenhum, o erro explica o motivo
    casos = [("/vazio/truncado", "limite de tokens"), ("/vazio/nulo", "vazia")]
    for rota, trecho in casos:
        monkeypatch.setitem(agents.PROVIDERS["openrouter"], "url", base + rota)
        with pytest.raises(agents.LLMError, match=trecho):
            agents.chat_conversa("openrouter", "k", "m", "CTX", [], "e ai?")

    monkeypatch.setitem(agents.PROVIDERS["anthropic"], "url", base + "/vazio/anthropic")
    with pytest.raises(agents.LLMError, match="limite de tokens"):
        agents.chat_conversa("anthropic", "k", "m", "CTX", [], "e ai?")

    monkeypatch.setitem(agents.PROVIDERS["google"], "url", base + "/vazio/gemini")
    with pytest.raises(agents.LLMError, match="SAFETY"):
        agents.chat_conversa("google", "k", "m", "CTX", [], "e ai?")


def test_cada_agente_fala_com_a_propria_especialidade(mock_llm):
    servidor, _ = mock_llm

    def sistema_de(agente):
        agents.chat_conversa("openrouter", "k", "m", "CTX", [], "e ai?", agente)
        return servidor.recebido["/openai"]["body"]["messages"][0]["content"]

    gestor = sistema_de("gestor")
    macro = sistema_de("macro")
    assert "gestor" in gestor.lower() and "posição" in gestor
    assert "economista" in macro.lower()
    assert gestor != macro

    # o quant conversa em texto e, quando propõe calibragem, fecha com o bloco
    # JSON que o painel transforma em botão — a decisão de aplicar é do usuário
    quant = sistema_de("premissas")
    assert "```json" in quant and "decisão de aplicar é do usuário" in quant

    # sem agente, é a mesa inteira falando junto
    mesa = sistema_de(None)
    assert "a mesa inteira" in mesa

    # a síntese recebe outro papel: fechar a reunião
    fim = sistema_de("sintese")
    assert "CONCLUSÃO" in fim and "fecha a reunião" in fim

    # o CONTEXTO entra em todos
    for s in (gestor, macro, quant, mesa, fim):
        assert "CTX" in s


def test_pergunta_de_sintese_carrega_o_que_a_mesa_disse():
    texto = agents.monta_pergunta_sintese("vale a pena?", [
        {"agente": "gestor", "nome": "Agente Gestor", "texto": "comprar"},
        {"agente": "macro", "nome": "Agente Macro", "texto": "juro caindo"},
        {"agente": "equity", "nome": "Vazio", "texto": "   "},
    ])
    assert "vale a pena?" in texto
    assert "--- Agente Gestor ---" in texto and "comprar" in texto
    assert "--- Agente Macro ---" in texto and "juro caindo" in texto
    # quem não respondeu não entra na conclusão
    assert "Vazio" not in texto


def test_conversa_trunca_historico_longo(mock_llm):
    servidor, _ = mock_llm
    hist = [{"role": "user", "content": f"m{i}"} for i in range(40)]
    agents.chat_conversa("openrouter", "sk-x", "m", "CTX", hist, "agora")
    corpo = servidor.recebido["/openai"]["body"]
    turnos = corpo["messages"][1:]
    assert len(turnos) == 13                     # 12 do histórico + a pergunta
    assert turnos[0]["content"] == "m28"


def test_parser_de_premissas_do_agente_quant():
    texto = (
        "Analisando o momento:\n```json\n"
        '{"premissas": {"rf": 0.148, "beta": 1.1, "growth": [0.09, 0.08, 0.07, 0.06, 0.05],'
        ' "g_terminal": 0.04}, "justificativa": "ok", "confianca": "media"}\n```\n'
    )
    proposta = agents.parse_assumption_json(texto)
    assert proposta["premissas"]["beta"] == 1.1
    assert len(proposta["premissas"]["growth"]) == 5
    assert agents.parse_assumption_json("sem json nenhum") is None
    assert agents.parse_assumption_json("") is None


def test_contexto_dos_agentes_nao_estoura_com_dados_faltando():
    payload = {"fundamentals": {"name": "X", "ticker": "XPTO3", "sector": "VAREJO",
                                "financial": False, "last_year": 2025,
                                "base": {}, "indicadores": {}, "series": {}, "years": []},
               "market": {"perf": {}}, "multiples": {}, "score": {}}
    texto = agents.build_context(payload, {}, {}, {})
    assert "XPTO3" in texto
    assert "sem dado" in texto


def test_todos_os_agentes_tem_prompt_e_descricao():
    for key, agente in agents.AGENTS.items():
        assert agente["system"].strip()
        assert agente["label"] and agente["desc"] and agente["icon"]
        assert "português" in agente["system"].lower()
    assert {a["key"] for a in agents.agent_list()} == set(agents.AGENTS)
    assert {p["key"] for p in agents.provider_list()} == set(agents.PROVIDERS)


# ---------------------------------------------------------------------------
# ETFs e BDRs
# ---------------------------------------------------------------------------

def test_universo_bdr_sem_duplicatas_e_setores_validos():
    from finlab.backend import bdrs
    tickers = [b.ticker for b in bdrs.UNIVERSE]
    assert len(tickers) == len(set(tickers))
    for b in bdrs.UNIVERSE:
        assert b.sector in bdrs.SECTORS, b.ticker
        assert b.us_ticker, b.ticker


def test_bdr_peers_do_mesmo_setor():
    from finlab.backend import bdrs
    pares = bdrs.peers("AAPL34")
    assert pares and all(p.sector == "TECHNOLOGY" for p in pares)
    assert all(p.ticker != "AAPL34" for p in pares)


def test_bancos_de_bdr_marcados():
    from finlab.backend import bdrs
    assert bdrs.get("JPMC34").bank
    assert not bdrs.get("VISA34").bank      # rede de pagamento, balanço corporativo
    assert not bdrs.get("AAPL34").bank


def test_fundamentos_de_bdr_a_partir_de_modulos_mockados():
    """Payload no formato Yahoo/BRAPI vira a mesma estrutura dos fundamentos CVM."""
    from finlab.backend import bdrs

    def stmt(ano, campos):
        base = {"endDate": {"fmt": f"{ano}-09-30"}}
        base.update({k: {"raw": v} for k, v in campos.items()})
        return base

    mod = {
        "incomeStatementHistory": {"incomeStatementHistory": [
            stmt(2025, {"totalRevenue": 400e9, "grossProfit": 180e9,
                        "ebit": 120e9, "netIncome": 100e9}),
            stmt(2024, {"totalRevenue": 380e9, "grossProfit": 170e9,
                        "ebit": 114e9, "netIncome": 95e9}),
            stmt(2023, {"totalRevenue": 360e9, "grossProfit": 160e9,
                        "ebit": 108e9, "netIncome": 90e9}),
            stmt(2022, {"totalRevenue": 340e9, "grossProfit": 150e9,
                        "ebit": 102e9, "netIncome": 85e9}),
        ]},
        "balanceSheetHistory": {"balanceSheetStatements": [
            stmt(2025, {"totalStockholderEquity": 70e9, "totalAssets": 350e9,
                        "cash": 30e9, "shortTermInvestments": 30e9,
                        "shortLongTermDebt": 10e9, "longTermDebt": 90e9}),
            stmt(2024, {"totalStockholderEquity": 65e9, "totalAssets": 340e9,
                        "cash": 28e9, "shortTermInvestments": 30e9,
                        "shortLongTermDebt": 11e9, "longTermDebt": 95e9}),
        ]},
        "cashflowStatementHistory": {"cashflowStatements": [
            stmt(2025, {"totalCashFromOperatingActivities": 110e9,
                        "capitalExpenditures": -12e9, "depreciation": 11e9}),
            stmt(2024, {"totalCashFromOperatingActivities": 105e9,
                        "capitalExpenditures": -11e9, "depreciation": 10e9}),
        ]},
        "financialData": {"financialCurrency": "USD"},
    }

    bdr = bdrs.get("AAPL34")
    fund = bdrs.fundamentals_from_modules(bdr, mod)

    assert fund["currency"] == "USD"
    assert fund["years"] == [2022, 2023, 2024, 2025]
    assert fund["last_year"] == 2025
    base = fund["base"]
    assert base["receita"] == pytest.approx(400e9)
    assert base["ebitda"] == pytest.approx(120e9 + 11e9)
    assert base["divida_bruta"] == pytest.approx(100e9)
    assert base["divida_liquida"] == pytest.approx(100e9 - 60e9)
    assert base["fcl"] == pytest.approx(110e9 - 12e9)
    ind = fund["indicadores"]
    assert ind["roe"] == pytest.approx(100e9 / 70e9)
    assert ind["cagr_receita_3a"] == pytest.approx((400 / 340) ** (1 / 3) - 1, rel=1e-6)
    # dá para pontuar com esses indicadores
    sc = scoring.score(ind, fund["financial"])
    assert sc["total"] is not None and 0 <= sc["total"] <= 100


def test_fundamentos_de_bdr_sem_modulos_ficam_vazios_e_sinalizados():
    from finlab.backend import bdrs
    fund = bdrs.fundamentals_from_modules(bdrs.get("MSFT34"), None)
    assert fund["years"] == []
    assert fund["bdr"] is True
    prem = valuation.bdr_assumptions(fund, {"price": 80.0}, {}, None)
    assert prem["aplicavel"] is False
    assert "BRAPI" in prem["motivo_nao_aplicavel"]


def test_valuation_de_bdr_converte_upside_para_preco_por_bdr(monkeypatch):
    """shares sintético: equity/shares = preço_BDR × (1+upside)."""
    from finlab.backend import b3data, bdrs

    monkeypatch.setattr(b3data, "usdbrl", lambda: 5.0)

    fund = {"sector": "TECHNOLOGY", "financial": False, "bdr": True, "currency": "USD",
            "years": [2023, 2024, 2025],
            "base": {"divida_bruta": 100e9, "divida_liquida": 40e9, "caixa": 60e9},
            "indicadores": {"cagr_receita_3a": 0.06},
            "series": {"fcl": [90e9, 95e9, 100e9], "ebit": [110e9, 115e9, 120e9],
                       "ebitda": [120e9, 126e9, 133e9]}}
    quote = {"marketCap": 5.0 * 2000e9}   # US$ 2 tri em BRL
    prem = valuation.bdr_assumptions(fund, {"price": 80.0}, {}, quote)

    assert prem["aplicavel"] is True
    assert prem["mcap_usd"] == pytest.approx(2000e9)
    # identidade da conversão
    assert prem["shares"] == pytest.approx(2000e9 / 80.0)
    equity_hipotetico = 2400e9   # +20% sobre o mcap
    preco_justo = equity_hipotetico / prem["shares"]
    assert preco_justo == pytest.approx(80.0 * 1.20)
    assert prem["g_terminal"] < prem["wacc"]


def test_universo_etf_completo_e_categorizado(monkeypatch):
    from finlab.backend import b3data, etfs as met

    monkeypatch.setattr(b3data, "etf_listing", lambda: [
        {"ticker": "BOVA11", "nome": "ISHARES IBOVESPA FUNDO DE ÍNDICE",
         "categoria_b3": "ETF Renda Variável"},
        {"ticker": "XPTO11", "nome": "GESTORA MSCI GLOBAL FUNDO DE ÍNDICE",
         "categoria_b3": "ETF Renda Variável"},
        {"ticker": "BOL5", "nome": "PRODUTO ESTRANHO", "categoria_b3": ""},
    ])
    monkeypatch.setattr(b3data, "registry_for", lambda nome: {"pl": 1e9, "pl_data": "2026-07-01",
                                                              "situacao": "Em Funcionamento Normal",
                                                              "gestor": None, "administrador": None,
                                                              "inicio": None})
    uni = met.universe()
    tickers = {e["ticker"] for e in uni}
    assert "BOVA11" in tickers
    assert "XPTO11" in tickers
    assert "BOL5" not in tickers            # código fora do padrão XXXX11 sai
    assert "HASH11" in tickers              # cripto entra pela lista extra

    bova = next(e for e in uni if e["ticker"] == "BOVA11")
    assert bova["categoria"] == "INDICES_BR"
    assert bova["curado"] and bova["taxa_adm"] == 0.10
    xpto = next(e for e in uni if e["ticker"] == "XPTO11")
    assert xpto["categoria"] == "INTERNACIONAL"   # heurística de nome
    assert not xpto["curado"]
    hash11 = next(e for e in uni if e["ticker"] == "HASH11")
    assert hash11["categoria"] == "CRIPTO"


def test_faixas_de_liquidez():
    from finlab.backend import etfs as met
    assert met.liquidity_band(None) == "sem negócios"
    assert met.liquidity_band(200e6) == "muito alta"
    assert met.liquidity_band(20e6) == "alta"
    assert met.liquidity_band(2e6) == "média"
    assert met.liquidity_band(200e3) == "baixa"
    assert met.liquidity_band(5e3) == "muito baixa"


def test_toda_meta_curada_de_etf_aponta_categoria_valida():
    from finlab.backend import etfs as met
    for ticker, meta in met.ETF_META.items():
        assert meta["cat"] in met.CATEGORIES, ticker
        assert meta["tese"], ticker
        if meta["taxa_adm"] is not None:
            assert 0 < meta["taxa_adm"] < 3, ticker


# ---------------------------------------------------------------------------
# Fundamentos de BDR via Yahoo Finance
# ---------------------------------------------------------------------------

def _yahoo_raw_exemplo():
    def anos(vals):
        return {2022 + i: v for i, v in enumerate(vals)}
    return {
        "income": {
            "receita": anos([340e9, 360e9, 380e9, 400e9]),
            "lucro_bruto": anos([150e9, 160e9, 170e9, 180e9]),
            "ebit": anos([102e9, 108e9, 114e9, 120e9]),
            "ebitda": anos([113e9, 119e9, 125e9, 131e9]),
            "lucro_liquido": anos([85e9, 90e9, 95e9, 100e9]),
        },
        "balance": {
            "patrimonio_liquido": anos([60e9, 62e9, 65e9, 70e9]),
            "ativo_total": anos([330e9, 335e9, 340e9, 350e9]),
            "caixa_total": anos([55e9, 58e9, 58e9, 60e9]),
            "divida_bruta": anos([105e9, 104e9, 106e9, 100e9]),
        },
        "cashflow": {
            "fco": anos([95e9, 100e9, 105e9, 110e9]),
            "capex": anos([-10e9, -10e9, -11e9, -12e9]),
            "depreciacao": anos([10e9, 10e9, 10e9, 11e9]),
        },
        "info": {"marketCap": 3.0e12, "beta": 1.2, "dividendYield": 0.44,
                 "currentPrice": 210.0, "targetMeanPrice": 250.0,
                 "targetHighPrice": 300.0, "targetLowPrice": 180.0,
                 "numberOfAnalystOpinions": 40, "recommendationKey": "buy",
                 "financialCurrency": "USD"},
    }


def test_fundamentos_de_bdr_via_yahoo():
    from finlab.backend import bdrs
    fund = bdrs.fundamentals_from_yahoo(bdrs.get("AAPL34"), _yahoo_raw_exemplo())
    assert fund["fonte"] == "Yahoo Finance"
    assert fund["years"] == [2022, 2023, 2024, 2025]
    base = fund["base"]
    assert base["receita"] == pytest.approx(400e9)
    assert base["ebitda"] == pytest.approx(131e9)          # linha EBITDA direto
    assert base["divida_liquida"] == pytest.approx(40e9)   # 100 − 60 (caixa consolidado)
    assert base["fcl"] == pytest.approx(98e9)
    sc = scoring.score(fund["indicadores"], fund["financial"])
    assert sc["total"] is not None


def test_yahoo_ebitda_derivado_quando_linha_falta():
    from finlab.backend import bdrs
    raw = _yahoo_raw_exemplo()
    del raw["income"]["ebitda"]
    fund = bdrs.fundamentals_from_yahoo(bdrs.get("AAPL34"), raw)
    # EBIT 120 + D&A 11
    assert fund["base"]["ebitda"] == pytest.approx(131e9)


def test_yahoo_banco_nao_ganha_ebitda_nem_divida_liquida():
    from finlab.backend import bdrs
    fund = bdrs.fundamentals_from_yahoo(bdrs.get("JPMC34"), _yahoo_raw_exemplo())
    assert fund["financial"] is True
    assert all(v is None for v in fund["series"]["ebitda"])
    assert fund["indicadores"]["nd_ebitda"] is None


def test_yahoo_dy_heuristica_de_escala():
    from finlab.backend import bdrs
    assert bdrs.yahoo_dividend_yield({"dividendYield": 0.0044}) == pytest.approx(0.0044)
    assert bdrs.yahoo_dividend_yield({"dividendYield": 0.44}) == pytest.approx(0.0044)
    assert bdrs.yahoo_dividend_yield({"dividendYield": None}) is None
    assert bdrs.yahoo_dividend_yield({}) is None


def test_orquestrador_prefere_yahoo_e_cai_para_brapi(monkeypatch):
    from finlab.backend import bdrs
    monkeypatch.setattr(bdrs, "yahoo_raw", lambda b: _yahoo_raw_exemplo())
    bundle = bdrs.fetch_fundamentals("AAPL34")
    assert bundle["fonte"] == "Yahoo Finance"
    assert bundle["info"]["marketCap"] == pytest.approx(3.0e12)

    monkeypatch.setattr(bdrs, "yahoo_raw", lambda b: None)
    monkeypatch.setattr(bdrs, "raw_modules", lambda t: None)
    bundle = bdrs.fetch_fundamentals("AAPL34")
    assert bundle["fonte"] is None
    assert bundle["fund"]["years"] == []


def test_bdr_assumptions_usa_mcap_do_yahoo_direto(monkeypatch):
    from finlab.backend import b3data, bdrs
    monkeypatch.setattr(b3data, "usdbrl", lambda: 5.0)
    fund = bdrs.fundamentals_from_yahoo(bdrs.get("AAPL34"), _yahoo_raw_exemplo())
    prem = valuation.bdr_assumptions(fund, {"price": 87.15}, {}, None,
                                     yahoo_info=_yahoo_raw_exemplo()["info"])
    assert prem["mcap_usd"] == pytest.approx(3.0e12)
    assert prem["mcap_fonte"] == "Yahoo Finance"
    assert prem["beta"] == pytest.approx(1.2)
    assert prem["beta_source"] == "Yahoo Finance"
    assert prem["aplicavel"] is True
    # identidade: equity == mcap → preço justo == preço do BDR
    assert 3.0e12 / prem["shares"] == pytest.approx(87.15)


def test_consenso_de_bdr_convertido_para_reais_por_bdr():
    from finlab.backend.app import _consenso_bdr
    cons = _consenso_bdr(_yahoo_raw_exemplo()["info"], 87.15)
    # alvo médio 250 sobre preço atual 210 → mesmo upside aplicado ao BDR
    assert cons["alvo_medio"] == pytest.approx(round(250 / 210 * 87.15, 2))
    assert cons["alvo_alto"] == pytest.approx(round(300 / 210 * 87.15, 2))
    assert cons["analistas"] == 40
    assert "Yahoo" in cons["fonte"]
    assert _consenso_bdr({}, 87.15) == {}
    assert _consenso_bdr(_yahoo_raw_exemplo()["info"], None) == {}


def test_df_to_plain_converte_dataframe_do_yfinance():
    import pandas as pd
    from finlab.backend import bdrs
    df = pd.DataFrame(
        {pd.Timestamp("2025-09-30"): [400e9, 100e9],
         pd.Timestamp("2024-09-30"): [380e9, float("nan")]},
        index=["Total Revenue", "Net Income"],
    )
    out = bdrs._df_to_plain(df, bdrs._Y_INCOME)
    assert out["receita"] == {2025: 400e9, 2024: 380e9}
    assert out["lucro_liquido"] == {2025: 100e9}   # NaN descartado
    assert bdrs._df_to_plain(None, bdrs._Y_INCOME) == {}


# ---------------------------------------------------------------------------
# Contexto dos agentes por tipo de ativo
# ---------------------------------------------------------------------------

def _payload_min(**over):
    fund = {"name": "X", "ticker": "XPTO3", "sector": "VAREJO", "financial": False,
            "last_year": 2025, "base": {"receita": 400e9}, "indicadores": {},
            "series": {}, "years": []}
    fund.update(over)
    return {"fundamentals": fund, "market": {"perf": {}}, "multiples": {}, "score": {}}


def test_contexto_de_acao_br_cita_cvm_e_reais():
    texto = agents.build_context(_payload_min(), {}, {}, {})
    assert "DFP" in texto and "CVM" in texto
    assert "R$ 400,00 bi" in texto
    assert "BDR" not in texto


def test_contexto_de_bdr_declara_origem_moeda_e_cambio():
    texto = agents.build_context(
        _payload_min(ticker="AAPL34", name="Apple", bdr=True, us_ticker="AAPL",
                     currency="USD", fonte="Yahoo Finance"), {}, {}, {})
    assert "BDR" in texto
    assert "Yahoo Finance" in texto
    assert "AAPL" in texto
    # grandezas contábeis na moeda de reporte
    assert "US$ 400,00 bi" in texto
    # e o aviso de que preço está em reais
    assert "REAIS por BDR" in texto
    assert "CVM" not in texto


def test_contexto_respeita_moeda_diferente_de_dolar():
    texto = agents.build_context(
        _payload_min(bdr=True, currency="EUR", us_ticker="ASML"), {}, {}, {})
    assert "EUR 400,00 bi" in texto


def test_prompt_do_macro_cobre_acao_br_e_bdr():
    sistema = agents.AGENTS["macro"]["system"]
    assert "BDR" in sistema
    assert "Tesouro americano" in sistema
    assert "Selic" in sistema


def test_regras_comuns_nao_prometem_cvm_para_todo_ativo():
    for agente in agents.AGENTS.values():
        assert "ORIGEM DOS DADOS" in agente["system"]


# ---------------------------------------------------------------------------
# Série trimestral (ITR)
# ---------------------------------------------------------------------------

RE_DS = "Receita de Venda de Bens e/ou Serviços"
LU_DS = "Lucro/Prejuízo Consolidado do Período"


def _linha_itr(fim, ini, conta, ds, valor, ordem="ÚLTIMO"):
    import pandas as pd

    return {"CD_CVM": "009512", "DENOM_CIA": "X", "CNPJ_CIA": "x",
            "DT_FIM_EXERC": pd.Timestamp(fim), "DT_INI_EXERC": pd.Timestamp(ini),
            "ORDEM_EXERC": ordem, "ANO_REFER": pd.Timestamp(fim).year,
            "CD_CONTA": conta, "DS_CONTA": ds, "VL_CONTA_AJUSTADO": valor}


def _monta_itr(tmp_path, linhas, anual=None):
    """Grava um ITR (e opcionalmente a DFP) sintéticos e devolve os pontos."""
    import pandas as pd

    pd.DataFrame(linhas).to_parquet(tmp_path / "dre_itr.parquet", index=False)
    if anual:
        pd.DataFrame([
            {"CD_CVM": "009512", "DENOM_CIA": "X", "CNPJ_CIA": "x",
             "DT_FIM_EXERC": pd.Timestamp(f"{ano}-12-31"), "ANO_REFER": ano,
             "CD_CONTA": conta, "DS_CONTA": ds, "VL_CONTA_AJUSTADO": v}
            for ano, conta, ds, v in anual
        ]).to_parquet(tmp_path / "dre_dfp.parquet", index=False)
    cvm.limpar_cache()
    return {p["rotulo"]: p for p in cvm.quarterly_series("009512")["pontos"]}


def test_serie_trimestral_desacumula_o_itr(tmp_path, monkeypatch):
    """A DRE do ITR vem acumulada no exercício. Plotar o acumulado como se
    fosse trimestre isolado desenha uma receita que só sobe — errado com cara
    de certo. Aqui: acumulado 100/220/360 vira 100/120/140."""
    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        linhas = []
        for fim, acc in (("2025-03-31", 100.0), ("2025-06-30", 220.0), ("2025-09-30", 360.0)):
            linhas += [_linha_itr(fim, "2025-01-01", "3.01", RE_DS, acc),
                       _linha_itr(fim, "2025-01-01", "3.11", LU_DS, acc / 10)]
        pontos = _monta_itr(tmp_path, linhas,
                            anual=[(2025, "3.01", RE_DS, 500.0), (2025, "3.11", LU_DS, 50.0)])

        assert pontos["1T25"]["receita"] == 100.0
        assert pontos["2T25"]["receita"] == 120.0
        assert pontos["3T25"]["receita"] == 140.0
        # o 4T não existe no ITR: sai do exercício fechado menos o acumulado
        assert pontos["4T25"]["receita"] == 140.0
        assert pontos["4T25"]["derivado"] is True
        assert pontos["1T25"]["derivado"] is False
        # validação forte: o LTM que fecha o exercício tem de bater com o anual
        assert pontos["4T25"]["receita_ltm"] == 500.0
        assert pontos["4T25"]["lucro_liquido_ltm"] == 50.0
    finally:
        cvm.limpar_cache()


def test_serie_trimestral_descarta_janela_avulsa_e_comparativo(tmp_path, monkeypatch):
    """O CSV do ITR traz, para a mesma data-fim, o acumulado e o trimestre
    avulso; e repete períodos antigos como exercício comparativo. Confundir
    qualquer um dos dois com o acumulado corrompe a diferença."""
    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        linhas = [
            _linha_itr("2025-03-31", "2025-01-01", "3.01", RE_DS, 100.0),
            _linha_itr("2025-06-30", "2025-01-01", "3.01", RE_DS, 220.0),
            # janela avulsa do 2T (abril–junho): valor diferente do acumulado
            _linha_itr("2025-06-30", "2025-04-01", "3.01", RE_DS, 777.0),
            # comparativo do ano anterior, possivelmente reapresentado
            _linha_itr("2024-03-31", "2024-01-01", "3.01", RE_DS, 999.0, ordem="PENÚLTIMO"),
        ]
        pontos = _monta_itr(tmp_path, linhas)

        assert pontos["2T25"]["receita"] == 120.0     # 220 − 100, não 777
        assert "1T24" not in pontos                   # o comparativo não vira ponto
    finally:
        cvm.limpar_cache()


def test_quarto_trimestre_so_sai_quando_o_terceiro_fechou(tmp_path, monkeypatch):
    """Se a empresa só publicou o 1T, anual − acumulado seriam nove meses
    empilhados num "4T". Melhor não desenhar do que desenhar errado."""
    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        pontos = _monta_itr(
            tmp_path,
            [_linha_itr("2025-03-31", "2025-01-01", "3.01", RE_DS, 100.0)],
            anual=[(2025, "3.01", RE_DS, 500.0)])

        assert list(pontos) == ["1T25"]
        assert pontos["1T25"]["receita"] == 100.0
    finally:
        cvm.limpar_cache()


def test_ltm_nao_soma_trimestres_com_buraco(tmp_path, monkeypatch):
    """Quatro pontos na série não são necessariamente quatro trimestres
    seguidos. Com um ano faltando, o LTM tem de ficar vazio em vez de somar
    períodos distantes e chamar isso de "últimos 12 meses"."""
    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        linhas = []
        for ano, acc in ((2022, (100.0, 220.0)), (2025, (130.0, 260.0))):
            for k, (fim_m, v) in enumerate(zip(("03-31", "06-30"), acc)):
                linhas.append(_linha_itr(f"{ano}-{fim_m}", f"{ano}-01-01", "3.01", RE_DS, v))
        pontos = _monta_itr(tmp_path, linhas)

        assert len(pontos) == 4
        assert all(p["receita_ltm"] is None for p in pontos.values())
    finally:
        cvm.limpar_cache()


def test_painel_segue_anual_sem_itr(tmp_path, monkeypatch):
    """Sem os parquets do ITR o painel não pode quebrar — só não mostra o
    trimestral."""
    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        assert cvm.quarterly_series("009512") == {"pontos": [], "campos": []}
        assert cvm.quarterly_series("") == {"pontos": [], "campos": []}
    finally:
        cvm.limpar_cache()


# ---------------------------------------------------------------------------
# Classificação de regime
# ---------------------------------------------------------------------------

def _fund(anos, **series):
    """Fundamentals mínimo no formato de cvm.annual_series."""
    return {"years": list(anos), "series": {k: list(v) for k, v in series.items()}}


def test_regime_sem_dado_nao_vira_operacao_normal():
    """A regra 1 do parecer: sem base para classificar, o painel diz que não
    sabe. Chutar R0 é pior que calar — R0 é justamente a hipótese que autoriza
    usar a média histórica como fluxo-base."""
    from finlab.backend import regime

    curto = regime.classificar(_fund([2024, 2025], lucro_liquido=[10.0, 12.0]))
    assert curto["codigo"] is None
    assert "3 exercícios" in curto["motivo"]

    vazio = regime.classificar(_fund([2020, 2021, 2022, 2023, 2024],
                                     receita=[1.0] * 5))
    assert vazio["codigo"] is None
    assert regime.classificar(None)["codigo"] is None


def test_regime_exige_dois_exercicios_para_expansao():
    """Um ano de capex alto é troca de frota. Regime pede confirmação."""
    from finlab.backend import regime

    anos = [2021, 2022, 2023, 2024, 2025]
    base = dict(receita=[100.0] * 5, depreciacao=[5.0] * 5,
                lucro_liquido=[10.0] * 5, imobilizado=[50, 55, 60, 70, 85],
                patrimonio_liquido=[100.0] * 5)

    um_ano = regime.classificar(_fund(anos, capex=[-5, -5, -5, -5, -30], **base))
    assert um_ano["codigo"] == "R0"

    dois = regime.classificar(_fund(anos, capex=[-5, -5, -5, -30, -30], **base))
    assert dois["codigo"] == "R1"


def test_regime_nao_confunde_empresa_leve_em_ativo_com_expansao():
    """Numa incorporadora a depreciação é quase nada, então capex/depreciação
    dispara sem que exista expansão: a MRV aparecia com 'capex de 7,2× a
    depreciação' investindo 4,7% da receita. A materialidade contra a receita
    é o que separa os dois casos."""
    from finlab.backend import regime

    anos = [2021, 2022, 2023, 2024, 2025]
    base = dict(receita=[100.0] * 5, lucro_liquido=[10.0] * 5,
                patrimonio_liquido=[100.0] * 5, imobilizado=[10, 11, 12, 13, 14])

    # capex 5× a depreciação, mas só 4% da receita: não é regime de capex
    leve = regime.classificar(_fund(anos, capex=[-4.0] * 5, depreciacao=[0.8] * 5, **base))
    assert leve["codigo"] == "R0"

    # mesma razão, capex de 20% da receita: aí sim
    pesada = regime.classificar(_fund(anos, capex=[-20.0] * 5, depreciacao=[4.0] * 5, **base))
    assert pesada["codigo"] == "R1"


def test_regime_ignora_arrumacao_de_portfolio():
    """Toda companhia mexe no portfólio; regime é outra coisa. Um item de 79 M
    sobre lucro de 921 M (o caso Totvs) não é reestruturação."""
    from finlab.backend import regime

    anos = [2021, 2022, 2023, 2024, 2025]
    base = dict(receita=[5000.0] * 5, capex=[-50.0] * 5, depreciacao=[40.0] * 5,
                patrimonio_liquido=[3000.0] * 5, lucro_liquido=[921.0] * 5)

    pequeno = regime.classificar(_fund(anos, descontinuadas=[0, 0, 0, 0, 79.0], **base))
    assert pequeno["codigo"] == "R0"

    grande = regime.classificar(_fund(anos, descontinuadas=[0, 0, 0, 0, 900.0], **base))
    assert grande["codigo"] == "R4"


def test_regime_pega_patrimonio_negativo_mesmo_com_lucro():
    """A Azul voltou ao azul em 2025 por 0,12 bi carregando patrimônio de
    −29 bi. Sem este sinal saía classificada como operação normal."""
    from finlab.backend import regime

    anos = [2021, 2022, 2023, 2024, 2025]
    r = regime.classificar(_fund(
        anos, receita=[16, 17, 18, 19, 21], lucro_liquido=[-0.7, -0.7, -2.4, -9.1, 0.12],
        patrimonio_liquido=[-19, -19, -21, -30, -29],
        capex=[-1.0] * 5, depreciacao=[-1.0] * 5))
    assert r["codigo"] == "R3"
    assert "patrimônio líquido negativo" in r["evidencias"][0]["texto"]
    assert r["evidencias"][0]["estrutural"] is True


def test_regime_escolhe_principal_por_precedencia_e_guarda_modificador():
    """A realidade combina: o caso GPA é turnaround COM desinvestimento. O
    principal é o que mais destrói a base do modelo; o outro vira modificador
    em vez de sumir."""
    from finlab.backend import regime

    anos = [2021, 2022, 2023, 2024, 2025]
    r = regime.classificar(_fund(
        anos, receita=[100.0] * 5, lucro_liquido=[5, 5, 5, -20.0, -10.0],
        descontinuadas=[0, 0, 0, -30.0, -25.0], patrimonio_liquido=[50.0] * 5,
        capex=[-2.0] * 5, depreciacao=[-2.0] * 5))
    assert r["codigo"] == "R3"
    assert r["modificador"]["codigo"] == "R4"
    # a evidência do modificador vem junto: o usuário vê as duas leituras
    assert {e["regime"] for e in r["evidencias"]} == {"R3", "R4"}


def test_regime_toda_evidencia_tem_data_e_o_texto_diz_a_conta():
    """Evidência sem data é opinião. E o painel promete mostrar a origem de
    todo número — inclusive os que ele mesmo inferiu."""
    from finlab.backend import regime

    anos = [2021, 2022, 2023, 2024, 2025]
    r = regime.classificar(_fund(
        anos, receita=[100.0] * 5, lucro_liquido=[5.0] * 5, patrimonio_liquido=[50.0] * 5,
        capex=[-20.0] * 5, depreciacao=[-4.0] * 5, imobilizado=[10, 20, 30, 40, 50]))
    assert r["evidencias"]
    for e in r["evidencias"]:
        assert isinstance(e["exercicio"], int) and 1990 < e["exercicio"] < 2100
        assert e["texto"] and e["regime"] == r["codigo"]
    # a confiança nunca promete mais do que a leitura contábil sustenta
    assert r["confianca"] in ("baixa", "media")


def test_regime_nao_derruba_o_painel_com_serie_torta():
    """Série malformada é rotina em dado público. O classificador degrada
    para 'sem classificação', nunca levanta exceção."""
    from finlab.backend import regime

    anos = [2021, 2022, 2023, 2024, 2025]
    for torta in (
        _fund(anos, lucro_liquido=[None, None, None, None, None]),
        _fund(anos, lucro_liquido=["x", None, 3.0, None, 1.0], receita=[None] * 5),
        _fund(anos, fco=[1.0] * 5, capex=[0.0] * 5, depreciacao=[0.0] * 5,
              receita=[0.0] * 5, lucro_liquido=[0.0] * 5),
    ):
        r = regime.classificar(torta)
        assert r["codigo"] in (None, "R0", "R1", "R2", "R3", "R4", "R5")


# ---------------------------------------------------------------------------
# Regime → fluxo-base (item 2.6)
# ---------------------------------------------------------------------------

def _fcf(ultimo, media3):
    return {"ultimo": ultimo, "media3": media3, "historico": []}


def test_regime_r0_nao_mexe_na_base():
    """R0 é o único mundo em que a média de 3 anos já é honesta. Mexer nela
    ali seria trocar uma premissa boa por outra sem motivo."""
    from finlab.backend import valuation

    fund = _fund([2023, 2024, 2025], fco=[10.0] * 3, depreciacao=[3.0] * 3)
    assert valuation.base_por_regime(fund, _fcf(7.0, 6.0), {"codigo": "R0"}) is None
    assert valuation.base_por_regime(fund, _fcf(7.0, 6.0), None) is None
    assert valuation.base_por_regime(fund, _fcf(7.0, 6.0), {"codigo": None}) is None


def test_base_de_expansao_troca_capex_total_por_manutencao():
    """Em R1 o FCL é negativo por escolha. A base do ativo maduro é o caixa
    das operações menos o capex de manutenção — e a depreciação é o proxy."""
    from finlab.backend import valuation

    fund = _fund([2023, 2024, 2025], fco=[100.0] * 3, depreciacao=[30.0] * 3)
    b = valuation.base_por_regime(fund, _fcf(-5.0, -8.0), {"codigo": "R1"})
    assert b["modo"] == "maduro"
    assert b["valor"] == 70.0                      # 100 − 30
    # o sinal da depreciação não pode inverter a conta
    fund_neg = _fund([2023, 2024, 2025], fco=[100.0] * 3, depreciacao=[-30.0] * 3)
    assert valuation.base_por_regime(fund_neg, _fcf(-5.0, -8.0),
                                     {"codigo": "R1"})["valor"] == 70.0
    # o texto avisa que o proxy é otimista em negócio intensivo em capital
    assert "subestima" in b["porque"]


def test_base_de_expansao_desiste_sem_depreciacao():
    """Sem depreciação não há proxy de manutenção: melhor não propor base
    nenhuma do que inventar uma."""
    from finlab.backend import valuation

    fund = _fund([2023, 2024, 2025], fco=[100.0] * 3)
    assert valuation.base_por_regime(fund, _fcf(1.0, 1.0), {"codigo": "R1"}) is None


def test_regimes_de_ruptura_usam_o_exercicio_mais_recente():
    """Em desalavancagem, turnaround e desinvestimento a média de 3 anos
    descreve outra empresa. Cada um explica o próprio motivo."""
    from finlab.backend import valuation

    fund = _fund([2023, 2024, 2025], fco=[10.0] * 3, depreciacao=[3.0] * 3)
    for codigo, marca in (("R2", "credor"), ("R3", "turnaround"), ("R4", "soma das partes")):
        b = valuation.base_por_regime(fund, _fcf(9.0, 4.0), {"codigo": codigo})
        assert b["modo"] == "ultimo" and b["valor"] == 9.0, codigo
        assert marca in b["porque"], codigo


def test_r5_nao_muda_a_base():
    """Em integração de M&A o que perde sentido é a comparação com pares, não
    o fluxo-base."""
    from finlab.backend import valuation

    fund = _fund([2023, 2024, 2025], fco=[10.0] * 3, depreciacao=[3.0] * 3)
    assert valuation.base_por_regime(fund, _fcf(9.0, 4.0), {"codigo": "R5"}) is None


def test_toda_base_de_regime_carrega_a_conta_e_o_porque():
    """O ajuste invisível é o pecado nº 6 do parecer: se o painel troca a base,
    ele tem de dizer qual conta produziu o número e por que trocou."""
    from finlab.backend import valuation

    fund = _fund([2023, 2024, 2025], fco=[100.0] * 3, depreciacao=[30.0] * 3)
    for codigo in ("R1", "R2", "R3", "R4"):
        b = valuation.base_por_regime(fund, _fcf(9.0, 4.0), {"codigo": codigo})
        assert b["conta"] and b["porque"] and b["rotulo"], codigo
        assert b["modo"] in ("maduro", "ultimo"), codigo


def test_base_de_regime_nao_resgata_fluxo_irrelevante():
    """Trocar uma base negativa por outra que é praticamente zero não melhora
    nada: o DCF passa a existir, mas o preço justo vira função só da dívida.
    A MRV saía de −1,40 bi para +28 mi e o painel cuspia −404% de upside onde
    antes dizia, honestamente, que não dava para valorar."""
    from finlab.backend import valuation

    fund = _fund([2023, 2024, 2025], ebitda=[0.20e9, 0.61e9, 0.45e9])
    # conversão de 7% do EBITDA: fica com o padrão e a barreira faz o trabalho
    assert valuation.base_por_regime(fund, _fcf(0.03e9, -1.40e9), {"codigo": "R3"}) is None
    # conversão saudável: a troca vale
    b = valuation.base_por_regime(fund, _fcf(0.30e9, -1.40e9), {"codigo": "R3"})
    assert b and b["valor"] == 0.30e9


# ---------------------------------------------------------------------------
# Ano em curso (LTM) na tabela de demonstrações
# ---------------------------------------------------------------------------

def test_ltm_soma_fluxo_e_nao_soma_saldo(tmp_path, monkeypatch):
    """As duas naturezas de conta: fluxo soma 12 meses, saldo é o do balanço
    mais recente. Somar quatro trimestres de patrimônio líquido seria absurdo,
    e é o tipo de erro que passa despercebido numa tabela bonita."""
    import pandas as pd

    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        # DRE trimestral acumulada: 4 trimestres isolados de 100 cada
        dre = []
        for ano, accs in ((2024, (100.0, 200.0, 300.0)), (2025, (100.0, 200.0, 300.0))):
            for fim, acc in zip((f"{ano}-03-31", f"{ano}-06-30", f"{ano}-09-30"), accs):
                dre.append(_linha_itr(fim, f"{ano}-01-01", "3.01", RE_DS, acc))
        pd.DataFrame(dre).to_parquet(tmp_path / "dre_itr.parquet", index=False)
        # anual fecha 2024 em 400 -> 4T24 isolado = 100
        pd.DataFrame([{
            "CD_CVM": "009512", "DENOM_CIA": "X", "CNPJ_CIA": "x",
            "DT_FIM_EXERC": pd.Timestamp("2024-12-31"), "ANO_REFER": 2024,
            "CD_CONTA": "3.01", "DS_CONTA": RE_DS, "VL_CONTA_AJUSTADO": 400.0,
        }]).to_parquet(tmp_path / "dre_dfp.parquet", index=False)
        # balanço: saldo cresce a cada trimestre; o LTM tem de pegar o último
        bpp = [{"CD_CVM": "009512", "DENOM_CIA": "X", "CNPJ_CIA": "x",
                "DT_INI_EXERC": pd.Timestamp("2025-01-01"),
                "DT_FIM_EXERC": pd.Timestamp(f), "ORDEM_EXERC": "ÚLTIMO",
                "ANO_REFER": 2025, "CD_CONTA": "2.03",
                "DS_CONTA": "Patrimônio Líquido Consolidado", "VL_CONTA_AJUSTADO": v}
               for f, v in (("2025-03-31", 900.0), ("2025-06-30", 950.0))]
        pd.DataFrame(bpp).to_parquet(tmp_path / "bpp_itr.parquet", index=False)
        cvm.limpar_cache()

        l = cvm.ltm_series("009512")
        assert l["fim"] == "2025-09-30"
        # fluxo: 4T24 (100) + 1T25 + 2T25 + 3T25 (100 cada) = 400
        assert l["campos"]["receita"] == 400.0
        # saldo: o do balanço mais recente, jamais a soma dos trimestres
        assert l["campos"]["patrimonio_liquido"] == 950.0
        assert "patrimonio_liquido" in l["saldos"]
    finally:
        cvm.limpar_cache()


def test_ltm_vazio_sem_itr(tmp_path, monkeypatch):
    """Sem ITR a coluna do ano em curso não aparece — e a tabela anual segue."""
    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        assert cvm.ltm_series("009512") == {}
        assert cvm.ltm_series("") == {}
    finally:
        cvm.limpar_cache()


# ---------------------------------------------------------------------------
# Radar de Contexto (busca ao vivo) e camada de momento
# ---------------------------------------------------------------------------

def test_busca_ao_vivo_so_no_provedor_que_tem_e_no_agente_que_pede(monkeypatch):
    """A busca externa é opt-in duplo: o provedor precisa suportar E o agente
    precisa pedir. Na forma nova (Agent Tools), pedir busca leva a chamada ao
    /v1/responses com `tools`; sem pedido — ou sem suporte — nada de tools."""
    chamadas = []

    def fake_post(url, **kw):
        chamadas.append((url, kw.get("json") or {}))
        if url.endswith("/responses"):
            return _RespFake(200, {"output_text": "ok"})
        return _RespFake(200, {"choices": [{"message": {"content": "ok"}}]})

    monkeypatch.setattr(agents.requests, "post", fake_post)

    agents.chat("xai", "k", "grok-4", "s", "u", buscar=True)
    assert chamadas[-1][0].endswith("/responses")
    assert chamadas[-1][1]["tools"] == agents.FERRAMENTAS_BUSCA

    agents.chat("xai", "k", "grok-4", "s", "u", buscar=False)
    assert chamadas[-1][0].endswith("/chat/completions")
    assert "tools" not in chamadas[-1][1]

    # OpenAI não tem busca ao vivo: pedir não muda nada.
    agents.chat("openai", "k", "gpt", "s", "u", buscar=True)
    assert chamadas[-1][0].endswith("/chat/completions")
    assert "tools" not in chamadas[-1][1] and "search_parameters" not in chamadas[-1][1]


def test_radar_abre_a_rodada_e_e_o_unico_com_busca():
    """Só um agente enxerga fora do painel — e é o que abre a rodada."""
    com_busca = [k for k, a in agents.AGENTS.items() if a.get("busca_ao_vivo")]
    assert com_busca == ["contexto"]
    assert agents.AGENTS["contexto"].get("abre_rodada") is True
    # o prompt exige as duas seções separadas, que é o ponto do desenho
    sistema = agents.AGENTS["contexto"]["system"]
    assert "FATO PUBLICADO" in sistema and "CONVERSA NÃO VERIFICADA" in sistema
    assert "DATA e LINK" in sistema
    # e proíbe o que estragaria a mesa
    assert "sentimento numérico" in sistema
    assert "Nunca opine sobre preço justo" in sistema


def test_momento_entra_no_contexto_com_evidencia_datada():
    """Sem isto o agente lia média de 3 anos de uma empresa em turnaround e
    tratava como run-rate — o erro que a classificação existe para evitar."""
    payload = {
        "fundamentals": {"last_year": 2025},
        "regime": {
            "codigo": "R3", "rotulo": "Turnaround", "confianca": "media",
            "modificador": {"codigo": "R4", "rotulo": "Reestruturação de portfólio"},
            "quebra": "o histórico deixa de ser âncora",
            "fluxo": "12 meses móveis do core",
            "evidencias": [{"exercicio": 2025, "texto": "prejuízo no exercício"}],
        },
        "trimestral": {"pontos": [
            {"rotulo": "4T24", "receita": 1e9, "lucro_liquido": 1e8, "derivado": True},
            {"rotulo": "1T25", "receita": 1.1e9, "lucro_liquido": 1.2e8},
        ]},
        "ltm": {"fim": "2025-03-31", "campos": {"receita": 4e9, "lucro_liquido": 4e8, "fcl": 2e8}},
    }
    texto = "\n".join(agents._bloco_momento(payload))
    assert "R3 · Turnaround" in texto and "com R4" in texto
    assert "[2025] prejuízo no exercício" in texto
    assert "4T24" in texto and "derivado da DFP" in texto
    assert "A mesa enxerga a contabilidade até 2025-03-31" in texto
    # e diz ao modelo o que ele NÃO sabe
    assert "SÓ contábil" in texto


def test_sem_regime_o_contexto_manda_nao_presumir_normal():
    """A regra 1 do parecer também vale para o que o agente lê."""
    texto = "\n".join(agents._bloco_momento(
        {"fundamentals": {"last_year": 2025},
         "regime": {"codigo": None, "motivo": "só 2 exercícios"}}))
    assert "Sem classificação" in texto
    assert "Não presuma operação normal" in texto


def test_bloco_de_momento_cala_quando_nao_ha_dado():
    """Cabeçalho vazio é convite para o modelo preencher sozinho."""
    assert agents._bloco_momento({}) == []
    assert agents._bloco_momento({"fundamentals": {}}) == []


def test_a_mesa_delibera_em_ondas_e_so_quem_fecha_le_os_outros():
    """Quatro leituras paralelas que nunca se cruzam não são uma mesa — são
    quatro monólogos, e sobra para o usuário achar onde elas discordam."""
    por_chave = {a["key"]: a for a in agents.agent_list()}

    # o Radar abre; o corpo da mesa vem junto; Cético e Moderador fecham nessa ordem
    assert por_chave["contexto"]["abre_rodada"] is True
    assert por_chave["cetico"]["ordem"] == 1
    assert por_chave["moderador"]["ordem"] == 2
    assert por_chave["equity"]["ordem"] == 0 and por_chave["gestor"]["ordem"] == 0

    # só quem fecha recebe as falas: dar o blackboard ao corpo da mesa faria
    # os quatro se ecoarem em vez de opinarem de forma independente
    leem = {k for k, a in por_chave.items() if a["le_a_mesa"]}
    assert leem == {"cetico", "moderador"}


def test_cetico_ataca_afirmacao_e_pode_nao_ter_o_que_contestar():
    """Cético que inventa disputa para parecer útil é pior que cético nenhum."""
    s = agents.AGENTS["cetico"]["system"]
    assert "NÃO produz tese própria" in s
    assert "nada a contestar" in s
    assert "Não invente disputa" in s
    # a prioridade certa: número fora do contexto é o erro mais grave
    assert "Número que não está no CONTEXTO" in s
    assert "Conversa tratada como fato" in s
    assert "Média histórica usada fora do regime" in s


def test_moderador_mapeia_disputa_em_vez_de_fabricar_consenso():
    """O parecer pede mapa de convergência/disputa NO LUGAR da síntese de
    consenso: média de opiniões esconde exatamente o que interessa."""
    s = agents.AGENTS["moderador"]["system"]
    assert "NÃO é um sintetizador de consenso" in s
    assert "CONVERGÊNCIA" in s and "DISPUTA" in s and "O QUE A MESA NÃO SABE" in s
    # a parte acionável: o que resolveria a discordância
    assert "o que decidiria" in s
    assert "Nunca produza recomendação" in s


def test_contexto_poe_o_estavel_antes_do_volatil():
    """Cache de prefixo só paga se o começo do prompt for byte a byte igual
    entre chamadas. Sete agentes leem a mesma empresa, e arrastar um slider
    deve revalidar só a cauda — não o dossiê inteiro."""
    fund = {"name": "X", "ticker": "X", "sector": "S", "last_year": 2025,
            "series": {}, "years": [], "base": {}, "indicadores": {}}
    payload = {"fundamentals": fund, "market": {}, "multiples": {}, "score": {},
               "regime": {"codigo": "R3", "rotulo": "Turnaround", "confianca": "media",
                          "quebra": "q", "fluxo": "f", "evidencias": []}}

    ctx = agents.build_context(payload, {"rf": 0.14}, {}, {})
    titulos = [l for l in ctx.splitlines() if l and not l.startswith(" ")]
    pos = {t: i for i, t in enumerate(titulos)}
    corte = next(i for t, i in pos.items() if "muda a cada ajuste" in t)

    # o dossiê da empresa fica todo antes do corte
    for estavel in ("HISTÓRICO", "FUNDAMENTOS (último exercício)"):
        assert pos[estavel] < corte, estavel
    assert next(i for t, i in pos.items() if t.startswith("MOMENTO")) < corte
    # e o que o slider mexe, todo depois
    assert pos["PREMISSAS ATUAIS DO PAINEL"] > corte
    assert pos["RESULTADO DO MODELO COM ESSAS PREMISSAS"] > corte

    # mudar premissa não pode alterar um único byte do prefixo estável
    outro = agents.build_context(payload, {"rf": 0.19, "beta": 2.0}, {"upside": 0.5}, {})
    marca = "--- daqui para baixo muda a cada ajuste de premissa ---"
    assert ctx.split(marca)[0] == outro.split(marca)[0]


def test_contexto_diz_sem_dado_em_vez_de_omitir():
    """Campo omitido é convite para o modelo preencher sozinho. Ausência tem
    de ser explícita — é a base do teste de abstenção do parecer 03."""
    ctx = agents.build_context(
        {"fundamentals": {"name": "X", "ticker": "X", "sector": "S", "series": {},
                          "years": [], "base": {}, "indicadores": {}},
         "market": {}, "multiples": {}, "score": {}},
        {}, {}, {})
    assert "sem dado" in ctx


def test_conjunto_de_abstencao_esta_bem_formado():
    """O eval de abstenção precisa de chave real, mas o conjunto em si é
    verificável offline — e um JSON quebrado só apareceria na hora errada."""
    import json
    from pathlib import Path

    dados = json.loads(
        (Path(__file__).parent / "golden" / "abstencao.json").read_text(encoding="utf-8"))
    assert dados["aceitas_como_abstencao"], "sem frases de abstenção não há teste"
    ids = [c["id"] for c in dados["casos"]]
    assert len(ids) == len(set(ids)), "id repetido no conjunto"
    for caso in dados["casos"]:
        assert caso["agente"] in agents.AGENTS, caso["id"]
        assert caso["pergunta"].endswith("?"), caso["id"]
        assert caso.get("porque"), f"{caso['id']} sem justificativa do caso"
    # os buracos que o plano ainda não fechou têm de estar cobertos
    assert {"guidance", "call", "fato_relevante"} <= set(ids)


# ---------------------------------------------------------------------------
# Índice IPE (o que a companhia comunicou)
# ---------------------------------------------------------------------------

def test_ipe_degrada_sem_o_indice(tmp_path, monkeypatch):
    """Sem o parquet do IPE o painel segue — como segue sem o ITR."""
    from finlab.backend import ipe

    monkeypatch.setattr(ipe, "CVM_PROCESSED_DIR", tmp_path)
    ipe._indice.cache_clear()
    try:
        assert ipe.disponivel() is False
        assert ipe.cobertura() is None
        assert ipe.documentos("009512") == {"docs": [], "cobertura": None, "total": 0}
        assert ipe.documentos("") == {"docs": [], "cobertura": None, "total": 0}
    finally:
        ipe._indice.cache_clear()


def test_ipe_casa_o_codigo_com_e_sem_zero_a_esquerda(tmp_path, monkeypatch):
    """O CD_CVM do universo vem zero-preenchido ("009512") e o do IPE nem
    sempre. Errar isso faria toda companhia aparecer sem documento nenhum."""
    import pandas as pd
    from finlab.backend import ipe

    monkeypatch.setattr(ipe, "CVM_PROCESSED_DIR", tmp_path)
    ipe._indice.cache_clear()
    try:
        pd.DataFrame([{
            "Codigo_CVM": "9512", "Categoria": "Fato Relevante", "Tipo": "",
            "Assunto": "Alienação de controlada", "Data_Entrega": pd.Timestamp("2026-07-01"),
            "Link_Download": "https://exemplo/1",
        }]).to_parquet(tmp_path / "ipe.parquet", index=False)
        ipe._indice.cache_clear()

        for chave in ("009512", "9512", " 009512 "):
            out = ipe.documentos(chave)
            assert out["total"] == 1, chave
            assert out["docs"][0]["link"] == "https://exemplo/1"
        assert ipe.documentos("000001")["total"] == 0
    finally:
        ipe._indice.cache_clear()


def test_ipe_prioriza_o_que_muda_tese_e_ordena_do_mais_novo(tmp_path, monkeypatch):
    """O IPE tem dezenas de categorias — assembleia, aviso aos acionistas,
    política de negociação. Mostrar tudo afogaria o fato relevante."""
    import pandas as pd
    from finlab.backend import ipe

    monkeypatch.setattr(ipe, "CVM_PROCESSED_DIR", tmp_path)
    ipe._indice.cache_clear()
    try:
        linhas = [
            ("Assembleia", "2026-07-30", "Edital de convocação"),
            ("Fato Relevante", "2026-07-01", "Venda de ativo"),
            ("Política de Negociação", "2026-06-20", "Atualização"),
            ("Comunicado ao Mercado", "2026-07-10", "Esclarecimento"),
        ]
        pd.DataFrame([{
            "Codigo_CVM": "9512", "Categoria": c, "Tipo": "", "Assunto": a,
            "Data_Entrega": pd.Timestamp(d), "Link_Download": "https://exemplo",
        } for c, d, a in linhas]).to_parquet(tmp_path / "ipe.parquet", index=False)
        ipe._indice.cache_clear()

        docs = ipe.documentos("9512")["docs"]
        cats = [d["categoria"] for d in docs]
        assert "Assembleia" not in cats and "Política de Negociação" not in cats
        # do mais novo para o mais velho: o que aconteceu por último é o que importa
        assert [d["data"] for d in docs] == ["2026-07-10", "2026-07-01"]
    finally:
        ipe._indice.cache_clear()


def test_ipe_entra_no_contexto_como_titulo_e_avisa_que_nao_leu_o_pdf():
    """A tentação óbvia é o agente inventar o conteúdo a partir do título."""
    texto = "\n".join(agents._bloco_momento({
        "fundamentals": {"last_year": 2025},
        "ipe": {"docs": [{"data": "2026-07-01", "categoria": "Fato Relevante",
                          "tipo": "", "assunto": "Venda de controlada"}]},
    }))
    assert "[2026-07-01] Fato Relevante: Venda de controlada" in texto
    assert "não pode afirmar o que está escrito dentro dele" in texto


# ---------------------------------------------------------------------------
# Conteúdo dos documentos (2.2/2.3): índice FTS5, busca e citação validada
# ---------------------------------------------------------------------------

def _pdf_minimo(texto: str) -> bytes:
    """Um PDF de verdade, com uma página e o texto pedido, sem dependência.

    Os offsets do xref são calculados, não chutados — pypdf valida a
    estrutura, e é justamente a extração real que o teste quer exercitar.
    """
    conteudo = f"BT /F1 11 Tf 40 700 Td ({texto}) Tj ET".encode("latin-1", "replace")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(conteudo), conteudo),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    saida = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, corpo in enumerate(objetos, start=1):
        offsets.append(len(saida))
        saida += b"%d 0 obj\n%s\nendobj\n" % (i, corpo)
    inicio_xref = len(saida)
    saida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    for off in offsets:
        saida += b"%010d 00000 n \n" % off
    saida += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF"
              % (len(objetos) + 1, inicio_xref))
    return bytes(saida)


def _importar_ipe_docs():
    raiz = Path(__file__).resolve().parents[2] / "valuation_cvm"
    sys.path.insert(0, str(raiz))
    try:
        from src import ipe_docs  # noqa: E402
        return ipe_docs
    finally:
        sys.path.pop(0)


def _montar_indice(tmp_path, registros):
    """docs.sqlite de teste, escrito pelas MESMAS funções do pipeline."""
    ipe_docs = _importar_ipe_docs()
    db = tmp_path / "docs.sqlite"
    original = ipe_docs.DB_PATH
    ipe_docs.DB_PATH = db
    try:
        con = ipe_docs._abrir_db()
        for meta, trechos in registros:
            ipe_docs._gravar_documento(con, meta, trechos, "ok" if trechos else "sem_texto")
        con.close()
    finally:
        ipe_docs.DB_PATH = original
    return db


_META_RESIA = {
    "protocolo": "P1", "cd_cvm": "9512", "categoria": "Fato Relevante",
    "tipo": "", "assunto": "Venda da Resia", "data_entrega": "2026-06-20",
    "data_referencia": None, "link": "https://rad.cvm.gov.br/ENET/doc1.pdf",
}
_META_DIV = {
    "protocolo": "P2", "cd_cvm": "9512", "categoria": "Comunicado ao Mercado",
    "tipo": "", "assunto": "Dividendos", "data_entrega": "2026-07-05",
    "data_referencia": None, "link": "https://rad.cvm.gov.br/ENET/doc2.pdf",
}


def test_busca_nos_documentos_devolve_data_e_link_sempre(tmp_path, monkeypatch):
    """O coração do 2.3: trecho sem data e sem link não pode existir.

    E a busca precisa ser insensível a acento — quem digita "aquisicao"
    tem de achar "aquisição".
    """
    from finlab.backend import docs as bdocs

    db = _montar_indice(tmp_path, [
        (_META_RESIA, ["A companhia comunica a venda da operação Resia nos EUA "
                       "por US$ 800 milhões, com efeito no 3T.",
                       "A aquisição de terrenos fica suspensa até a conclusão."]),
        (_META_DIV, ["Distribuição de dividendos intermediários aprovada."]),
    ])
    monkeypatch.setattr(bdocs, "DB_PATH", db)

    achados = bdocs.search("009512", "o que aconteceu com a Resia?")
    assert achados, "a busca tinha de achar o fato relevante"
    assert achados[0]["data"] == "2026-06-20"
    assert achados[0]["link"].startswith("https://rad.cvm.gov.br/")
    assert "Resia" in achados[0]["trecho"]

    # sem acento acha com acento
    sem_acento = bdocs.search("9512", "aquisicao de terrenos")
    assert any("aquisição" in t["trecho"] for t in sem_acento)

    # recentes: mais novo primeiro, só o primeiro trecho de cada documento
    rec = bdocs.recentes("9512", n=2)
    assert [r["data"] for r in rec] == ["2026-07-05", "2026-06-20"]

    st = bdocs.stats("9512")
    assert st["disponivel"] is True and st["documentos"] == 2
    assert st["ultimo"] == "2026-07-05"


def test_documentos_degradam_para_vazio_sem_indice(tmp_path, monkeypatch):
    from finlab.backend import docs as bdocs
    monkeypatch.setattr(bdocs, "DB_PATH", tmp_path / "nao-existe.sqlite")
    assert bdocs.available() is False
    assert bdocs.search("9512", "resia") == []
    assert bdocs.recentes("9512") == []
    assert bdocs.stats()["disponivel"] is False


def test_bloco_de_contexto_data_antes_do_trecho_e_abstencao_no_vazio():
    from finlab.backend import docs as bdocs

    bloco = bdocs.bloco_contexto([{
        "data": "2026-06-20", "categoria": "Fato Relevante",
        "assunto": "Venda da Resia", "link": "https://rad.cvm.gov.br/ENET/doc1.pdf",
        "protocolo": "P1", "trecho": "venda da operação Resia",
    }])
    # o ID e a data aparecem antes do texto do trecho, e o link vai junto
    assert bloco.index("[doc P1 · 2026-06-20]") < bloco.index("venda da operação")
    assert "https://rad.cvm.gov.br/ENET/doc1.pdf" in bloco

    vazio = bdocs.bloco_contexto([])
    assert "nenhum trecho recuperado" in vazio
    assert "não há documento recuperado" in vazio


def test_citacao_de_link_inventado_e_marcada_em_codigo():
    """Parecer 03 §5: validação em código, não em prompt. Link do RAD que o
    modelo não recebeu sai marcado; o recebido passa limpo."""
    from finlab.backend import docs as bdocs

    trechos = [{"link": "https://rad.cvm.gov.br/ENET/doc1.pdf"}]
    texto = ("A venda consta em https://rad.cvm.gov.br/ENET/doc1.pdf e o guidance "
             "em https://rad.cvm.gov.br/ENET/inventado.pdf.")
    saida = bdocs.validar_citacoes(texto, trechos)
    assert "doc1.pdf e o guidance" in saida
    assert "inventado.pdf. ⚠[link não recuperado nesta consulta]" in saida
    # link de outro domínio não é da alçada desta validação
    assert bdocs.validar_citacoes("veja https://exemplo.com/x", []) == "veja https://exemplo.com/x"


def test_etapa_de_documentos_extrai_indexa_e_e_incremental(tmp_path, monkeypatch):
    """A etapa inteira do pipeline com um PDF de verdade — e sem rede: o
    arquivo já está no lugar, então nem download nem Crawl-Delay acontecem.
    Na segunda rodada o protocolo já indexado não é tocado."""
    import pandas as pd
    ipe_docs = _importar_ipe_docs()

    ipe = pd.DataFrame([{
        "Codigo_CVM": "9512", "Categoria": "Fato Relevante", "Tipo": "",
        "Assunto": "Venda da Resia", "Data_Entrega": "2026-06-20",
        "Data_Referencia": "2026-06-20", "Protocolo_Entrega": "PX1",
        "Link_Download": "https://rad.cvm.gov.br/ENET/doc1.pdf", "Versao": "1",
    }])
    (tmp_path / "processed").mkdir()
    ipe.to_parquet(tmp_path / "processed" / "ipe.parquet", index=False)

    monkeypatch.setattr(ipe_docs, "PROCESSED_DIR", tmp_path / "processed")
    monkeypatch.setattr(ipe_docs, "DOCS_DIR", tmp_path / "docs")
    monkeypatch.setattr(ipe_docs, "DB_PATH", tmp_path / "processed" / "docs.sqlite")
    monkeypatch.setattr(ipe_docs, "_codigos_do_universo", lambda: {"9512"})
    monkeypatch.setattr(ipe_docs.time, "sleep",
                        lambda s: (_ for _ in ()).throw(AssertionError("dormiu sem baixar")))

    destino = tmp_path / "docs" / "9512" / "PX1.pdf"
    destino.parent.mkdir(parents=True)
    destino.write_bytes(_pdf_minimo("Venda da operacao Resia por US$ 800 milhoes"))

    placar = ipe_docs.indexar(meses=6000, por_empresa=5)
    assert placar == {"baixados": 1, "pulados": 0, "falhas": 0}

    from finlab.backend import docs as bdocs
    monkeypatch.setattr(bdocs, "DB_PATH", tmp_path / "processed" / "docs.sqlite")
    achados = bdocs.search("9512", "Resia")
    assert achados and "Resia" in achados[0]["trecho"]
    assert achados[0]["data"] == "2026-06-20"

    # segunda rodada: nada a baixar, nada refeito
    placar2 = ipe_docs.indexar(meses=6000, por_empresa=5)
    assert placar2 == {"baixados": 0, "pulados": 1, "falhas": 0}


def test_selecao_respeita_categoria_universo_janela_e_teto():
    import pandas as pd
    ipe_docs = _importar_ipe_docs()

    hoje = pd.Timestamp.today()
    linhas = []
    for i in range(10):
        linhas.append({"Codigo_CVM": "9512", "Categoria": "Fato Relevante",
                       "Data_Entrega": hoje - pd.Timedelta(days=i * 30),
                       "Protocolo_Entrega": f"A{i}", "Link_Download": "https://x/a.pdf"})
    linhas.append({"Codigo_CVM": "9512", "Categoria": "Assembleia",
                   "Data_Entrega": hoje, "Protocolo_Entrega": "IRRELEV",
                   "Link_Download": "https://x/b.pdf"})
    linhas.append({"Codigo_CVM": "777777", "Categoria": "Fato Relevante",
                   "Data_Entrega": hoje, "Protocolo_Entrega": "FORA",
                   "Link_Download": "https://x/c.pdf"})
    linhas.append({"Codigo_CVM": "9512", "Categoria": "Fato Relevante",
                   "Data_Entrega": hoje - pd.Timedelta(days=900),
                   "Protocolo_Entrega": "VELHO", "Link_Download": "https://x/d.pdf"})

    sel = ipe_docs._selecionar(pd.DataFrame(linhas), meses=24, por_empresa=4,
                               universo={"9512"})
    protocolos = list(sel["Protocolo_Entrega"])
    assert len(protocolos) == 4                      # teto por empresa
    assert "IRRELEV" not in protocolos               # categoria fora da lista
    assert "FORA" not in protocolos                  # empresa fora do universo
    assert "VELHO" not in protocolos                 # fora da janela
    assert protocolos == sorted(protocolos, key=lambda p: int(p[1:]))  # mais novos


def test_corte_em_paragrafos_com_rabicho_juntado():
    ipe_docs = _importar_ipe_docs()
    paragrafo = "x" * 500
    texto = "\n\n".join([paragrafo, paragrafo, paragrafo, "fim curto"])
    trechos = ipe_docs._cortar(texto)
    assert all(len(t) <= ipe_docs.CHUNK_ALVO + 600 for t in trechos)
    # o rabicho curto não vira trecho próprio
    assert trechos[-1].endswith("fim curto") and len(trechos[-1]) > len("fim curto")
    assert ipe_docs._cortar("") == []


def test_extrair_pdf_do_anexo_e_bloco_rotulado():
    """O anexo do chat: extração real, teto de tamanho e rótulo honesto."""
    from finlab.backend import docs as bdocs

    anexo = bdocs.extrair_pdf(_pdf_minimo("Release de resultados do 2T26"))
    assert "Release de resultados" in anexo["texto"]
    assert anexo["paginas"] == 1 and anexo["truncado"] is False

    with pytest.raises(ValueError, match="não parece ser um PDF"):
        bdocs.extrair_pdf(b"isto e um txt")
    with pytest.raises(ValueError, match="15 MB"):
        bdocs.extrair_pdf(b"%PDF" + b"0" * (bdocs.ANEXO_MAX_BYTES + 1))

    bloco = bdocs.bloco_anexo("release.pdf", {"texto": "conteudo", "truncado": True})
    assert "DOCUMENTO ENVIADO PELO USUÁRIO — release.pdf" in bloco
    assert "truncado" in bloco
    # o rótulo separa material do usuário de fonte oficial
    assert "NÃO" in bloco and "demonstrações da CVM" in bloco


def test_chat_em_streaming_deltas_uso_e_queda_para_inteiro(monkeypatch):
    """3.6: deltas enquanto o modelo escreve, tokens no fechamento — e o
    provedor que recusar o stream cai para a chamada inteira sem falhar."""
    class RespStream:
        status_code = 200

        def iter_lines(self, decode_unicode=False):
            yield ": keep-alive"
            yield 'data: {"choices":[{"delta":{"content":"fundamentos "}}]}'
            yield 'data: {"choices":[{"delta":{"content":"sólidos"}}]}'
            yield 'data: {"choices":[],"usage":{"prompt_tokens":900,"completion_tokens":42}}'
            yield "data: [DONE]"

        def close(self):
            pass

    monkeypatch.setattr(agents.requests, "post", lambda *a, **kw: RespStream())
    eventos = list(agents.chat_conversa_stream("openrouter", "k", "m", "CTX", [], "oi"))
    assert [e.get("delta") for e in eventos[:-1]] == ["fundamentos ", "sólidos"]
    fim = eventos[-1]
    assert fim["fim"] is True and fim["texto"] == "fundamentos sólidos"
    assert fim["uso"] == {"entrada": 900, "saida": 42}

    # provedor recusa o stream (400) → resolve inteiro, num delta único
    chamadas = []

    def post_sem_stream(url, **kw):
        chamadas.append(bool((kw.get("json") or {}).get("stream")))
        if (kw.get("json") or {}).get("stream"):
            return _RespFake(400, {"error": "stream não suportado"})
        return _RespFake(200, {"choices": [{"message": {"content": "inteiro"}}]})

    monkeypatch.setattr(agents.requests, "post", post_sem_stream)
    eventos = list(agents.chat_conversa_stream("openrouter", "k", "m", "CTX", [], "oi"))
    assert chamadas == [True, False]
    assert eventos == [{"delta": "inteiro"},
                       {"fim": True, "texto": "inteiro", "uso": None}]

    # busca ao vivo não passa pelo stream: vai direto ao /v1/responses
    def post_responses(url, **kw):
        assert url.endswith("/responses"), "busca tinha de ir ao responses"
        return _RespFake(200, {"output_text": "radar"})

    monkeypatch.setattr(agents.requests, "post", post_responses)
    eventos = list(agents.chat_conversa_stream("xai", "k", "m", "CTX", [], "oi",
                                               "contexto", buscar=True))
    assert eventos[-1]["texto"] == "radar"


def test_fato_com_doc_id_e_contestacao_verificavel():
    """3.4: o bloco carrega o ID de cada trecho e a regra fato × interpretação;
    (doc X) com ID fora do recuperado é marcado em código, não em prompt."""
    from finlab.backend import docs as bdocs

    trechos = [{"data": "2026-06-20", "categoria": "Fato Relevante",
                "assunto": "Venda", "link": "https://rad.cvm.gov.br/ENET/doc1.pdf",
                "protocolo": "F1", "trecho": "venda por US$ 800 milhões"}]
    bloco = bdocs.bloco_contexto(trechos)
    assert "[doc F1 · 2026-06-20]" in bloco
    assert "(doc ID)" in bloco and "INTERPRETAÇÃO" in bloco

    texto = ("A venda foi por US$ 800 milhões (doc F1). O follow-on está "
             "aprovado (doc F9). Minha leitura: a alavancagem cai.")
    saida = bdocs.validar_citacoes(texto, trechos)
    assert "(doc F1)." in saida and "(doc F1) ⚠" not in saida
    assert "(doc F9) ⚠[doc inexistente no recuperado]" in saida
    # interpretação sem etiqueta passa intocada
    assert "Minha leitura: a alavancagem cai." in saida

    # e o Cético é instruído a contestar pelo MESMO ID
    assert "(doc ID)" in agents.FECHAMENTO_DA_RODADA["cetico"]
    assert "doc ID" in agents.AGENTS["cetico"]["system"]


def test_reconciliador_devolve_proposta_so_do_quant_com_empresa_aberta(monkeypatch):
    """4.3: a proposta de premissas vira dado estruturado na resposta do chat —
    mas só quando quem falou é o Engenheiro de Premissas E há empresa aberta."""
    from finlab.backend import app as bapp

    resposta = ("Rf de 14% e beta 1,1 fazem mais sentido hoje.\n\n"
                "```json\n"
                '{"premissas": {"rf": 0.14, "beta": 1.1}, "confianca": "media"}\n'
                "```")
    monkeypatch.setattr(bapp.agents, "chat_conversa", lambda *a, **kw: resposta)

    corpo = {"slot": {"provider": "openrouter", "api_key": "k", "model": "m"},
             "ticker": "PETR4", "pergunta": "calibra o modelo",
             "agente": "premissas"}
    r = bapp.api_agent_chat(dict(corpo))
    assert r["proposta"]["premissas"] == {"rf": 0.14, "beta": 1.1}

    # outro agente com o mesmo texto: nada de proposta
    r2 = bapp.api_agent_chat(dict(corpo, agente="equity"))
    assert r2["proposta"] is None

    # sem ticker do universo: não há onde aplicar
    r3 = bapp.api_agent_chat(dict(corpo, ticker=""))
    assert r3["proposta"] is None


# ---------------------------------------------------------------------------
# Placar de promessas (4.2)
# ---------------------------------------------------------------------------

@pytest.fixture
def placar_limpo(tmp_path, monkeypatch):
    """Cada teste com o seu arquivo — o placar é estado em disco."""
    from finlab.backend import promessas as pr
    monkeypatch.setattr(pr, "ARQUIVO", tmp_path / "promessas.json")
    return pr


def test_promessa_versionada_nunca_perde_o_historico(placar_limpo):
    """O ponto do placar: promessa que muda de prazo duas vezes é o dado mais
    valioso aqui. Um UPDATE que sobrescrevesse apagaria exatamente isso."""
    pr = placar_limpo

    p = pr.registrar("MRVE3", {"texto": "desalavancar para 2,0x",
                               "prazo": "2026-12-31", "metrica": "DL/EBITDA"})
    assert p["estado"] == "aberta" and p["revisoes"] == 0

    pr.atualizar("MRVE3", p["id"], {"prazo": "2027-06-30",
                                    "nota": "adiado no call do 3T"})
    depois = pr.atualizar("MRVE3", p["id"], {"estado": "quebrada",
                                             "nota": "fechou 2027 em 2,8x"})

    assert depois["estado"] == "quebrada"
    assert depois["prazo"] == "2027-06-30"     # herdou o prazo revisado
    assert depois["texto"] == "desalavancar para 2,0x"   # texto nunca reenviado
    assert depois["revisoes"] == 2
    # o histórico inteiro sobreviveu, na ordem
    prazos = [v["prazo"] for v in depois["versoes"]]
    assert prazos == ["2026-12-31", "2027-06-30", "2027-06-30"]
    assert [v["estado"] for v in depois["versoes"]] == ["aberta", "aberta", "quebrada"]


def test_placar_conta_vencidas_e_so_calcula_taxa_com_resolvidas(placar_limpo):
    """Taxa de cumprimento sem nenhuma promessa resolvida seria 0% — uma
    mentira aritmética sobre uma gestão que ainda não foi cobrada."""
    from datetime import date
    pr = placar_limpo
    hoje = date(2026, 8, 6)

    pr.registrar("PETR4", {"texto": "capex de 20 bi", "prazo": "2026-01-31"})
    pr.registrar("PETR4", {"texto": "venda de refinaria", "prazo": "2027-12-31"})
    pr.registrar("PETR4", {"texto": "sem prazo declarado"})

    p = pr.placar("PETR4", hoje)
    assert p["total"] == 3 and p["aberta"] == 3
    assert p["vencidas"] == 1, "só a de 31/01 passou do prazo"
    assert p["taxa"] is None, "nada resolvido ainda"
    # vencida primeiro na lista: é o que exige ação
    assert p["itens"][0]["texto"] == "capex de 20 bi"

    alvo = p["itens"][0]["id"]
    pr.atualizar("PETR4", alvo, {"estado": "cumprida"})
    p2 = pr.placar("PETR4", hoje)
    assert p2["cumprida"] == 1 and p2["vencidas"] == 0
    assert p2["taxa"] == 1.0, "1 de 1 resolvida foi cumprida"


def test_promessas_isoladas_por_ticker_e_remocao(placar_limpo):
    pr = placar_limpo
    a = pr.registrar("VALE3", {"texto": "dividendo extraordinário"})
    pr.registrar("PETR4", {"texto": "outra empresa"})

    assert len(pr.listar("VALE3")) == 1
    assert len(pr.listar("PETR4")) == 1
    assert pr.listar("BBAS3") == []

    assert pr.remover("VALE3", a["id"]) is True
    assert pr.listar("VALE3") == []
    assert pr.remover("VALE3", a["id"]) is False       # já não existe
    assert len(pr.listar("PETR4")) == 1, "remover num ticker não toca no outro"

    with pytest.raises(ValueError):
        pr.registrar("VALE3", {"texto": "   "})        # texto é obrigatório
    with pytest.raises(ValueError):
        pr.registrar("VALE3", {"texto": "x", "estado": "talvez"})
    with pytest.raises(KeyError):
        pr.atualizar("VALE3", "inexistente", {"estado": "cumprida"})


def test_arquivo_corrompido_nao_derruba_o_placar(placar_limpo):
    pr = placar_limpo
    pr.ARQUIVO.parent.mkdir(parents=True, exist_ok=True)
    pr.ARQUIVO.write_text("{isto não é json", encoding="utf-8")
    assert pr.listar("PETR4") == []
    assert pr.placar("PETR4")["total"] == 0
    # e volta a gravar por cima, sem exigir intervenção
    assert pr.registrar("PETR4", {"texto": "recomeço"})["texto"] == "recomeço"


def test_placar_entra_no_contexto_da_mesa_com_a_proibicao_de_inventar():
    """Sem a proibição, o agente completa a lista com promessa plausível —
    que é a falha mais cara que este placar pode ter."""
    from datetime import date

    payload = {
        "fundamentals": {"last_year": 2025},
        "promessas": {
            "total": 2, "aberta": 1, "cumprida": 1, "quebrada": 0, "parcial": 0,
            "vencidas": 1, "taxa": 1.0,
            "itens": [
                {"texto": "desalavancar para 2,0x", "prazo": "2026-01-31",
                 "estado": "aberta", "vencida": True, "revisoes": 2,
                 "doc": "F1", "nota": "adiado duas vezes"},
                {"texto": "venda da Resia", "prazo": "2025-12-31",
                 "estado": "cumprida", "vencida": False, "revisoes": 0},
            ],
        },
    }
    texto = "\n".join(agents._bloco_momento(payload))
    assert "PLACAR DE PROMESSAS" in texto
    assert "[VENCIDA] prazo 2026-01-31 (doc F1) · replanejada 2x" in texto
    assert "adiado duas vezes" in texto
    assert "cumprimento 100%" in texto
    assert "NUNCA afirme cumprimento ou descumprimento de promessa que não esteja" in texto

    # sem promessa registrada, o bloco não aparece — silêncio em vez de
    # cabeçalho vazio, que o modelo tende a preencher sozinho
    vazio = "\n".join(agents._bloco_momento({"fundamentals": {"last_year": 2025},
                                             "promessas": {"total": 0, "itens": []}}))
    assert "PLACAR DE PROMESSAS" not in vazio


def test_promessa_extraida_e_filtrada_pelo_documento_recuperado():
    """A ponte documento → placar: o agente propõe, mas quem valida a
    procedência é o código. Promessa ancorada em doc inventado não chega
    à tela, e o link vem do índice, não do que o modelo escreveu."""
    from finlab.backend import docs as bdocs

    trechos = [{"protocolo": "F1", "data": "2026-06-20",
                "link": "https://rad.cvm.gov.br/ENET/doc1.pdf",
                "categoria": "Fato Relevante", "assunto": "Venda",
                "trecho": "..."}]
    texto = (
        "Encontrei dois compromissos.\n\n```json\n"
        '{"promessas": ['
        '{"texto": "desalavancar para 2,0x", "prazo": "2026-12-31", '
        '"metrica": "DL/EBITDA", "doc": "F1", "link": "https://mentira.com/x"},'
        '{"texto": "promessa de documento que não veio", "doc": "F9"},'
        '{"texto": "prazo impossível", "prazo": "2026-13-45", "doc": "F1"},'
        '{"texto": "", "doc": "F1"}'
        ']}\n```'
    )
    achadas = bdocs.promessas_propostas(texto, trechos)

    assert len(achadas) == 2, "só as ancoradas em F1 com texto"
    primeira = achadas[0]
    assert primeira["texto"] == "desalavancar para 2,0x"
    assert primeira["prazo"] == "2026-12-31"
    assert primeira["origem"] == "documento"
    # o link e a data vêm do ÍNDICE, não do que o modelo escreveu
    assert primeira["link"] == "https://rad.cvm.gov.br/ENET/doc1.pdf"
    assert primeira["data_origem"] == "2026-06-20"
    # data inválida é descartada, mas a promessa continua (sem prazo)
    assert achadas[1]["texto"] == "prazo impossível" and achadas[1]["prazo"] is None

    # sem bloco, sem documento, ou com JSON quebrado: lista vazia, sem exceção
    assert bdocs.promessas_propostas("só texto", trechos) == []
    assert bdocs.promessas_propostas(texto, []) == []
    assert bdocs.promessas_propostas('```json\n{"promessas": [ quebrado\n```', trechos) == []


def test_regra_de_extracao_so_entra_com_documento_no_contexto():
    """Ensinar o formato a quem não tem de onde extrair é convite a inventar."""
    com_doc = agents._sistema_da_conversa(
        "equity", "DOCUMENTOS DA EMPRESA (trechos oficiais recuperados)\n[doc F1 · 2026-01-01]")
    sem_doc = agents._sistema_da_conversa("equity", "MOMENTO DA EMPRESA\n  Regime: R0")

    assert "EXTRAIR PROMESSAS DA GESTÃO" in com_doc
    assert "AAAA-MM-DD" in com_doc and "não invente data" in com_doc
    assert "EXTRAIR PROMESSAS DA GESTÃO" not in sem_doc


# ---------------------------------------------------------------------------
# Transcrição de teleconferência (4.1)
# ---------------------------------------------------------------------------

_CALL = """WEBVTT

1
00:00:01.000 --> 00:00:05.000
Operador: Bom dia. Iniciamos a teleconferência do 2T26.

João Silva, CEO: O trimestre teve margem de 34%, recorde da companhia.

Maria Souza, CFO: A dívida líquida fechou em 2,8x EBITDA.

Operador: Iniciaremos agora a sessão de perguntas e respostas.

Operador: A próxima pergunta vem de Pedro Lima, do Banco X.

Pedro Lima: Qual o prazo para chegar em 2,0x? E o capex de 2027 muda?

João Silva, CEO: Esperamos 2,0x até o 4T27. O capex segue em 3 bilhões.

Operador: Próxima pergunta, de Ana Costa.

Ana Costa: A venda da unidade internacional continua no radar?

Maria Souza, CFO: Sim, mantemos o processo, sem prazo definido.
"""


def test_call_segmenta_em_pares_pergunta_resposta():
    """A unidade de recuperação de uma call é a TROCA inteira: quem pergunta
    sobre alavancagem precisa da pergunta e da resposta juntas — metade
    engana. Legenda .vtt e falas do operador não podem virar conteúdo."""
    from finlab.backend import calls

    s = calls.segmentar(_CALL)

    assert len(s["qa"]) == 2
    assert s["qa"][0]["quem_pergunta"] == "Pedro Lima"
    assert "2,0x" in s["qa"][0]["pergunta"] and "capex" in s["qa"][0]["pergunta"]
    assert s["qa"][0]["quem_resposta"] == "João Silva, CEO"
    assert "4T27" in s["qa"][0]["resposta"]
    assert s["qa"][1]["quem_pergunta"] == "Ana Costa"

    # apresentação: só os dois executivos; o operador não entra
    assert len(s["apresentacao"]) == 2
    assert not any("Operador" in b for b in s["apresentacao"])
    # nem o cabeçalho da legenda nem o timestamp sobrevivem
    juntos = "\n".join(s["trechos"])
    assert "WEBVTT" not in juntos and "00:00:01" not in juntos
    # o trecho do Q&A carrega a troca inteira
    qa = [t for t in s["trechos"] if t.startswith("[Q&A]")]
    assert len(qa) == 2 and "P (Pedro Lima)" in qa[0] and "R (João Silva" in qa[0]


def test_call_sem_rotulo_de_falante_e_sem_sessao_de_perguntas():
    """Transcrição crua (ASR sem diarização) não pode falhar — vira
    apresentação inteira em vez de exceção. E a marca de sessão de perguntas
    aparecendo sem que haja perguntas não pode engolir a call: melhor tudo
    como apresentação do que texto perdido."""
    from finlab.backend import calls

    s = calls.segmentar("Texto corrido, sem falantes rotulados nesta gravação. " * 5)
    assert s["qa"] == []
    assert s["apresentacao"] and s["trechos"]

    # a marca no meio da apresentação, sem nenhum "?" depois
    falso = calls.segmentar("Ana, CEO: Depois vem a sessão de perguntas e respostas. "
                            "Seguimos com o plano de capex.")
    assert falso["qa"] == []
    assert falso["apresentacao"], "o corte errado apagou a call inteira"
    assert "capex" in " ".join(falso["trechos"])
    assert calls.segmentar("")["trechos"] == []
    assert calls.segmentar("   \n\n  ")["trechos"] == []


def test_call_indexada_vira_documento_citavel_pela_mesa(tmp_path, monkeypatch):
    """O ponto do desenho: a call entra no MESMO índice, então herda de graça
    a recuperação com data e doc ID — e a mesa a cita como qualquer documento."""
    from finlab.backend import calls, docs as bdocs

    monkeypatch.setattr(bdocs, "DB_PATH", tmp_path / "docs.sqlite")

    r = calls.indexar("009512", "2026-08-05", _CALL, "Call do 2T26")
    assert r["qa"] == 2 and r["protocolo"] == "call-2026-08-05"

    achados = bdocs.search("9512", "prazo para chegar em 2,0x")
    assert achados, "a call não foi recuperada"
    assert achados[0]["data"] == "2026-08-05"
    assert achados[0]["protocolo"] == "call-2026-08-05"
    assert achados[0]["categoria"] == "Teleconferência"
    # a troca inteira veio junta
    assert "P (Pedro Lima)" in achados[0]["trecho"] and "4T27" in achados[0]["trecho"]

    # e o bloco de contexto a rotula com o ID, como os outros documentos
    bloco = bdocs.bloco_contexto(achados[:1])
    assert "[doc call-2026-08-05 · 2026-08-05]" in bloco

    # listar, regravar (substitui, não duplica) e remover
    assert [c["protocolo"] for c in calls.listar("009512")] == ["call-2026-08-05"]
    calls.indexar("009512", "2026-08-05", _CALL, "Call do 2T26 (corrigida)")
    assert len(calls.listar("9512")) == 1, "regravar duplicou a call"
    assert calls.listar("9512")[0]["titulo"] == "Call do 2T26 (corrigida)"

    assert calls.remover("9512", "call-2026-08-05") is True
    assert calls.listar("9512") == []
    assert calls.remover("9512", "call-2026-08-05") is False

    with pytest.raises(ValueError):
        calls.indexar("009512", "05/08/2026", _CALL)     # data fora do ISO
    with pytest.raises(ValueError):
        calls.indexar("009512", "2026-08-05", "   ")     # sem texto


def test_esquema_do_indice_nao_diverge_entre_app_e_pipeline(tmp_path):
    """As duas cópias do DDL (backend/docs.py e o pipeline) precisam descrever
    o mesmo índice — divergir faria a call e os PDFs viverem em tabelas
    incompatíveis, e o erro só apareceria em produção."""
    import sqlite3
    from finlab.backend import docs as bdocs
    ipe_docs = _importar_ipe_docs()

    def colunas(ddl, arquivo):
        con = sqlite3.connect(tmp_path / arquivo)
        con.executescript(ddl)
        fora = {}
        for tabela in ("documentos", "trechos"):
            fora[tabela] = [r[1] for r in con.execute(f"PRAGMA table_info({tabela})")]
        con.close()
        return fora

    assert colunas(bdocs.ESQUEMA, "a.sqlite") == colunas(ipe_docs.ESQUEMA, "b.sqlite")


def test_mesa_fecha_com_modelo_lendo_as_falas(monkeypatch):
    """A discussão vira um conjunto de premissas defensável: o quant recebe as
    falas da rodada e o fechamento exige a curva de crescimento — que é a
    projeção do fluxo — e a declaração de onde a mesa discordou."""
    from finlab.backend import app as bapp

    visto = {}

    def fake_chat(provider, api_key, model, contexto, historico, pergunta,
                  agente=None, buscar=False):
        visto["pergunta"] = pergunta
        visto["agente"] = agente
        return ("Adotei o lado do Cético na margem.\n\n```json\n"
                '{"premissas": {"rf": 0.14, "growth": [0.05,0.04,0.04,0.03,0.03]},'
                ' "confianca": "media"}\n```')

    monkeypatch.setattr(bapp.agents, "chat_conversa", fake_chat)

    r = bapp.api_agent_chat({
        "slot": {"provider": "openrouter", "api_key": "k", "model": "m"},
        "ticker": "PETR4", "pergunta": "vale a posição?", "agente": "premissas",
        "respostas": [{"agente": "equity", "nome": "Analista", "texto": "margem sobe"},
                      {"agente": "cetico", "nome": "Cético", "texto": "margem não sobe"}],
    })

    # as falas chegaram ao quant, com a instrução de fechamento
    assert "O QUE A MESA RESPONDEU" in visto["pergunta"]
    assert "margem não sobe" in visto["pergunta"]
    assert "curva de crescimento" in visto["pergunta"]
    assert "onde a mesa DISCORDOU" in visto["pergunta"]
    # e a proposta volta estruturada, com a curva
    assert r["proposta"]["premissas"]["growth"] == [0.05, 0.04, 0.04, 0.03, 0.03]

    # sem respostas, o quant continua respondendo a pergunta crua
    visto.clear()
    bapp.api_agent_chat({
        "slot": {"provider": "openrouter", "api_key": "k", "model": "m"},
        "ticker": "PETR4", "pergunta": "e o beta?", "agente": "premissas",
    })
    assert visto["pergunta"] == "e o beta?"


def test_epv_do_servidor_bate_com_o_motor_do_navegador():
    """O EPV existe em dois lugares — engine.js (por empresa, com sliders) e
    valuation.py (as 90 da tela principal). Divergir seria o painel dizer um
    número e a mesa outro sobre a MESMA empresa."""
    prem = {"ebit_normalizado": 8000.0, "wacc": 0.152, "tax": 0.34,
            "divida_liquida": 12000.0, "shares": 4196.0, "preco": 41.20}
    v = valuation.epv(prem)

    # a conta do engine.js, escrita à mão aqui: EBIT após imposto / WACC
    esperado_valor = 8000.0 * (1 - 0.34) / 0.152
    assert v["valor"] == pytest.approx(esperado_valor)
    assert v["equity"] == pytest.approx(esperado_valor - 12000.0)
    assert v["por_acao"] == pytest.approx((esperado_valor - 12000.0) / 4196.0)
    assert v["upside"] == pytest.approx(v["por_acao"] / 41.20 - 1)

    # dívida líquida ausente é zero, não erro (caixa líquido some no None)
    assert valuation.epv({"ebit_normalizado": 100.0, "wacc": 0.1,
                          "shares": 10.0})["equity"] == pytest.approx(660.0)
    # sem EBIT, sem WACC ou com WACC ≤ 0 não há valor — e nada explode
    assert valuation.epv({"wacc": 0.1})["por_acao"] is None
    assert valuation.epv({"ebit_normalizado": 100.0, "wacc": 0})["por_acao"] is None
    assert valuation.epv({})["valor"] is None
    # sem ações, o valor da firma existe mas o por-ação não
    semacoes = valuation.epv({"ebit_normalizado": 100.0, "wacc": 0.1, "tax": 0.0})
    assert semacoes["valor"] == pytest.approx(1000.0) and semacoes["por_acao"] is None


def test_linha_da_tela_principal_traz_epv_lpa_e_nao_cai_com_dado_torto(monkeypatch):
    """A linha do overview é o que a mesa lê: EPV para operacional, LPA para
    todo mundo, e nenhuma empresa com dado torto pode derrubar a tela."""
    from finlab.backend import app as bapp

    fund = {"financial": False, "base": {"lucro_liquido": 6840.0}}
    snap = {"shares_quote": 4196.0, "price": 41.20}
    monkeypatch.setattr(bapp.valuation, "assumptions", lambda *a, **kw: {
        "ebit_normalizado": 8230.0, "wacc": 0.152, "tax": 0.34,
        "divida_liquida": -3100.0, "shares": 4196.0, "preco": 41.20})

    linha = bapp._epv_da_linha(fund, snap, {}, None)
    assert linha["epv_por_acao"] is not None and linha["epv_upside"] is not None
    assert linha["lpa"] == pytest.approx(6840.0 / 4196.0, abs=1e-4)
    assert linha["wacc"] == 0.152

    # financeira: sem EPV por definição, mas COM LPA — é a métrica do setor
    banco = bapp._epv_da_linha({"financial": True, "base": {"lucro_liquido": 1000.0}},
                               {"shares_quote": 500.0}, {}, None)
    assert banco["epv_por_acao"] is None
    assert banco["lpa"] == pytest.approx(2.0)

    # premissa que explode não derruba a linha inteira
    def explode(*a, **kw):
        raise RuntimeError("dado torto")
    monkeypatch.setattr(bapp.valuation, "assumptions", explode)
    caido = bapp._epv_da_linha(fund, snap, {}, None)
    assert caido["epv_por_acao"] is None
    assert caido["lpa"] == pytest.approx(6840.0 / 4196.0, abs=1e-4), "o LPA sobrevive"


# ---------------------------------------------------------------------------
# DRE estruturada (redesenho · spec 4.3)
# ---------------------------------------------------------------------------

def test_dre_anual_monta_as_linhas_com_margens_e_cagr(tmp_path, monkeypatch):
    """A DRE de leitura: ordem contábil, margens derivadas da receita, CAGR só
    onde ele significa alguma coisa, e EBITDA = EBIT + |D&A| (a CVM não
    publica EBITDA como conta)."""
    import pandas as pd

    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        linhas = []
        for ano, receita, cpv, bruto, desp, ebit, resfin, ir, lucro in [
            (2024, 1000.0, -600.0, 400.0, -150.0, 250.0, -20.0, -60.0, 170.0),
            (2025, 1200.0, -700.0, 500.0, -180.0, 320.0, -25.0, -80.0, 215.0),
        ]:
            for conta, ds, v in [("3.01", "Receita de Venda de Bens", receita),
                                 ("3.02", "Custo dos Bens Vendidos", cpv),
                                 ("3.03", "Resultado Bruto", bruto),
                                 ("3.04", "Despesas/Receitas Operacionais", desp),
                                 ("3.05", "Resultado Antes do Resultado Financeiro", ebit),
                                 ("3.06", "Resultado Financeiro", resfin),
                                 ("3.08", "Imposto de Renda e Contribuição Social", ir),
                                 ("3.11", "Lucro/Prejuizo Consolidado do Periodo", lucro)]:
                linhas.append({"CD_CVM": "009512", "DENOM_CIA": "T", "CNPJ_CIA": "1",
                               "ANO_REFER": ano, "ORDEM_EXERC": "ÚLTIMO",
                               "DT_FIM_EXERC": f"{ano}-12-31",
                               "CD_CONTA": conta, "DS_CONTA": ds, "VL_CONTA_AJUSTADO": v})
        pd.DataFrame(linhas).to_parquet(tmp_path / "dre_dfp.parquet", index=False)
        # D&A vive na DFC, como ajuste do FCO
        pd.DataFrame([
            {"CD_CVM": "009512", "DENOM_CIA": "T", "CNPJ_CIA": "1", "ANO_REFER": ano,
             "ORDEM_EXERC": "ÚLTIMO", "DT_FIM_EXERC": f"{ano}-12-31",
             "CD_CONTA": "6.01.01.02", "DS_CONTA": "Depreciacao e Amortizacao",
             "VL_CONTA_AJUSTADO": v}
            for ano, v in [(2024, 50.0), (2025, 60.0)]
        ]).to_parquet(tmp_path / "dfc_mi_dfp.parquet", index=False)
        cvm.limpar_cache()

        d = cvm.dre_anual("009512")
        assert d["anos"] == [2024, 2025] and d["financial"] is False
        por = {l["chave"]: l for l in d["linhas"]}

        # ordem contábil preservada
        assert [l["chave"] for l in d["linhas"]][:4] == [
            "receita", "cpv", "lucro_bruto", "mg_bruta"]
        assert por["receita"]["valores"] == [1000.0, 1200.0]
        assert por["cpv"]["tipo"] == "deducao"
        assert por["lucro_liquido"]["tipo"] == "hero"

        # EBITDA derivado: EBIT + |D&A|
        assert por["ebitda"]["valores"] == [300.0, 380.0]
        assert por["da"]["valores"] == [-50.0, -60.0], "D&A entra negativa"

        # margens sobre a receita, na linha própria
        assert por["mg_bruta"]["valores"] == [pytest.approx(0.4), pytest.approx(0.4167, abs=1e-3)]
        assert por["mg_ebitda"]["valores"][1] == pytest.approx(380.0 / 1200.0)
        assert por["mg_liquida"]["tipo"] == "margem"

        # CAGR só nas linhas de resultado; margem e dedução não têm
        assert por["receita"]["cagr"] == pytest.approx(0.2)
        assert por["mg_bruta"]["cagr"] is None
        assert por["cpv"]["cagr"] is None
    finally:
        cvm.limpar_cache()


def test_dre_anual_de_financeira_cai_no_plano_reduzido(tmp_path, monkeypatch):
    """Banco não tem CPV nem EBITDA: mostrar as linhas vazias sugeriria que o
    dado faltou. E linha só de zeros (o IR da consolidada) não entra."""
    import pandas as pd

    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        pd.DataFrame([
            {"CD_CVM": "009512", "DENOM_CIA": "B", "CNPJ_CIA": "1", "ANO_REFER": 2025,
             "ORDEM_EXERC": "ÚLTIMO", "DT_FIM_EXERC": "2025-12-31",
             "CD_CONTA": c, "DS_CONTA": ds, "VL_CONTA_AJUSTADO": v}
            for c, ds, v in [
                ("3.01", "Receitas da Intermediação Financeira", 500.0),
                ("3.06", "Resultado Financeiro", -10.0),
                ("3.08", "Imposto de Renda e Contribuição Social", 0.0),
                ("3.11", "Lucro/Prejuizo Consolidado do Periodo", 90.0)]
        ]).to_parquet(tmp_path / "dre_dfp.parquet", index=False)
        cvm.limpar_cache()

        d = cvm.dre_anual("009512")
        chaves = [l["chave"] for l in d["linhas"]]
        assert d["financial"] is True
        assert "cpv" not in chaves and "ebitda" not in chaves and "da" not in chaves
        assert "receita" in chaves and "lucro_liquido" in chaves and "mg_liquida" in chaves
        assert "ir" not in chaves, "linha só de zeros não é informação"
    finally:
        cvm.limpar_cache()


def test_dre_trimestral_desacumula_soma_o_acumulado_e_compara_com_o_ano_anterior(
        tmp_path, monkeypatch):
    """Três invariantes: o ITR vem acumulado e sai isolado; a coluna do
    semestre é a soma dos trimestres; e o Δ a/a compara o mesmo trimestre —
    não o anterior, que a sazonalidade distorce."""
    import pandas as pd

    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        def linha(ano, mes_fim, acumulado, valor, conta="3.01",
                  ds="Receita de Venda de Bens"):
            return {"CD_CVM": "009512", "DENOM_CIA": "T", "CNPJ_CIA": "1",
                    "ANO_REFER": ano, "ORDEM_EXERC": "ÚLTIMO",
                    "DT_INI_EXERC": f"{ano}-01-01",
                    "DT_FIM_EXERC": f"{ano}-{mes_fim}", "CD_CONTA": conta,
                    "DS_CONTA": ds, "VL_CONTA_AJUSTADO": acumulado if acumulado else valor}

        itr = []
        for ano, (a1, a2) in [(2024, (100.0, 220.0)), (2025, (120.0, 260.0))]:
            for conta, ds, fator in [("3.01", "Receita de Venda de Bens", 1.0),
                                     ("3.11", "Lucro/Prejuizo Consolidado do Periodo", 0.2)]:
                itr.append(linha(ano, "03-31", a1 * fator, None, conta, ds))
                itr.append(linha(ano, "06-30", a2 * fator, None, conta, ds))
        pd.DataFrame(itr).to_parquet(tmp_path / "dre_itr.parquet", index=False)
        cvm.limpar_cache()

        t = cvm.dre_trimestral("009512")
        assert t["ano"] == 2025
        assert [c["rotulo"] for c in t["colunas"]] == ["1T25", "2T25", "1S25"]
        assert t["colunas"][2]["acumulado"] is True

        receita = next(l for l in t["linhas"] if l["chave"] == "receita")
        # acumulado 120 e 260 → trimestres isolados 120 e 140
        assert receita["valores"] == [120.0, 140.0, 260.0]
        assert receita["valores"][2] == receita["valores"][0] + receita["valores"][1]

        # Δ a/a por índice de trimestre: 1T25/1T24 e 2T25/2T24 (=140/120)
        assert receita["yoy"][0] == pytest.approx(120.0 / 100.0 - 1)
        assert receita["yoy"][1] == pytest.approx(140.0 / 120.0 - 1)
        assert receita["yoy"][2] == pytest.approx(260.0 / 220.0 - 1)

        # margem não tem Δ a/a em pontos percentuais nesta tabela
        margem = next(l for l in t["linhas"] if l["chave"] == "mg_liquida")
        assert all(y is None for y in margem["yoy"])
    finally:
        cvm.limpar_cache()


def test_dre_degrada_sem_dado(tmp_path, monkeypatch):
    monkeypatch.setattr(cvm, "CVM_PROCESSED_DIR", tmp_path)
    cvm.limpar_cache()
    try:
        assert cvm.dre_anual("009512")["linhas"] == []
        assert cvm.dre_trimestral("009512")["colunas"] == []
        assert cvm.dre_anual("")["anos"] == []
        assert cvm.dre_trimestral("")["linhas"] == []
    finally:
        cvm.limpar_cache()


def test_cagr_recusa_sinal_trocado():
    """Prejuízo virando lucro não tem taxa composta que signifique nada."""
    # a janela é a distância entre o primeiro e o último valor VÁLIDO —
    # série que só começa no meio não vira CAGR de 5 anos
    assert cvm._cagr([100.0, 121.0], 5) == pytest.approx(0.21)
    assert cvm._cagr([None, None, 100.0, 121.0], 3) == pytest.approx(0.21)
    assert cvm._cagr([100.0, 110.0, 121.0], 2) == pytest.approx(0.1)
    assert cvm._cagr([-50.0, 121.0], 1) is None
    assert cvm._cagr([100.0, -20.0], 1) is None
    assert cvm._cagr([None, 100.0], 1) is None
    assert cvm._cagr([], 5) is None


# ---------------------------------------------------------------------------
# Planilha DCF exportada (spec 4.1 do redesenho)
# ---------------------------------------------------------------------------

def _prem_planilha(**sobrepor):
    """Premissas mínimas no formato de valuation.assumptions — valores em R$
    (a planilha converte para milhões)."""
    prem = {
        "fcf_media3": 4.8e9, "fcf_ultimo": 5.6e9,
        "divida_liquida": -3.1e9,
        "shares": 4.196e9, "preco": 41.20,
        "growth": [0.12, 0.105, 0.09, 0.08, 0.07], "g_terminal": 0.05,
        "rf": 0.131, "rf_fonte": "ANBIMA · prefixado ~10 anos",
        "beta": 0.72, "erp": 0.045, "premio_extra": 0.0,
        "spread_credito": 0.015, "wd": 0.05, "tax": 0.20, "wacc": 0.16,
    }
    prem.update(sobrepor)
    return prem


def _abrir_planilha(blob):
    import io
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(blob))


def test_planilha_dcf_tem_formulas_vivas_e_nunca_resultado_pronto():
    """A regra central do endpoint: resultado é SEMPRE fórmula. Se a planilha
    trouxesse números calculados em Python, um bug de conta ficaria escondido
    atrás de um valor que parece certo."""
    from finlab.backend import xlsx_dcf

    blob = xlsx_dcf.planilha("WEGE3", _prem_planilha(),
                             {"last_year": 2025, "financial": False})
    wb = _abrir_planilha(blob)
    ws = wb["DCF"]

    # custo de capital e WACC, célula a célula como no arquivo de referência
    assert ws["B29"].value == "=B20+B21*B22+B23"
    assert ws["C30"].value == "=(C20+C24)*(1-C26)"
    assert ws["D31"].value == "=D29*(1-D25)+D30*D25"

    # projeção encadeada e valor presente
    assert ws["C34"].value == "=C13*(1+C14)"
    assert ws["C38"].value == "=C37*(1+C18)"
    assert ws["B41"].value == "=B34/(1+B$31)^1"
    assert ws["D45"].value == "=D38/(1+D$31)^5"
    assert ws["C46"].value == "=SUM(C41:C45)"

    # a guarda de Gordon: WACC ≤ g devolve "n/a", não um número sem sentido
    assert ws["C47"].value == '=IF(C31<=C19,"n/a",C38*(1+C19)/(C31-C19))'
    assert ws["C53"].value == '=IF(ISNUMBER(C51),C51/$B$9,"n/a")'
    assert ws["B54"].value == '=IF(ISNUMBER(B53),B53/$B$10-1,"n/a")'

    # nenhuma célula de resultado (linhas 29–54) veio como número pronto
    for linha in range(29, 55):
        for col in ("B", "C", "D"):
            v = ws[f"{col}{linha}"].value
            if v is not None:
                assert str(v).startswith("="), f"{col}{linha} veio calculado: {v}"


def test_planilha_dcf_dados_do_painel_e_tres_cenarios():
    from finlab.backend import xlsx_dcf

    blob = xlsx_dcf.planilha("WEGE3", _prem_planilha(),
                             {"last_year": 2025, "financial": False})
    ws = _abrir_planilha(blob)["DCF"]

    # dados do painel em R$ milhões, com a célula de fonte ao lado
    assert ws["B6"].value == 4800
    assert ws["B7"].value == 5600
    assert ws["B8"].value == -3100
    assert ws["B9"].value == 4196
    assert ws["B10"].value == 41.20
    assert "DFC da CVM 2023–2025" in str(ws["C6"].value)

    # três cenários: pessimista < mediana < otimista no crescimento e no g
    assert ws["B12"].value == "Pessimista"
    assert ws["C12"].value == "Mediana"
    assert ws["D12"].value == "Otimista"
    assert ws["C14"].value == pytest.approx(0.12)
    assert ws["B14"].value < ws["C14"].value < ws["D14"].value
    assert ws["B19"].value < ws["C19"].value < ws["D19"].value
    assert ws["C19"].value == pytest.approx(0.05)
    # beta pessimista é o MAIOR (mais desconto), como no arquivo de referência
    assert ws["B21"].value > ws["C21"].value > ws["D21"].value
    # FCL base igual nos três: o cenário mexe em crescimento e risco
    assert ws["B13"].value == ws["C13"].value == ws["D13"].value == 4800

    # células editáveis marcadas (fundo amarelo), dado do painel em cinza
    assert ws["C14"].fill.start_color.rgb == "FFFFFF00"
    assert ws["B6"].fill.start_color.rgb == "FFF2F2F2"
    assert ws["C14"].font.color.rgb == "FF0000FF"
    assert ws["A1"].font.name == "Arial"


def test_planilha_dcf_sensibilidade_referencia_os_fluxos_da_mediana():
    from finlab.backend import xlsx_dcf

    blob = xlsx_dcf.planilha("WEGE3", _prem_planilha(),
                             {"last_year": 2025, "financial": False})
    wb = _abrir_planilha(blob)
    assert "Sensibilidade" in wb.sheetnames
    ws = wb["Sensibilidade"]

    # grade centrada no painel: WACC 16% ± 3 p.p., g 5% ± 2 p.p.
    assert ws["A4"].value == "WACC \\ g"
    assert ws["B4"].value == pytest.approx(0.03)
    assert ws["F4"].value == pytest.approx(0.07)
    assert ws["A5"].value == pytest.approx(0.13)
    assert ws["A11"].value == pytest.approx(0.19)

    celula = str(ws["B5"].value)
    assert celula.startswith('=IF($A5<=B$4,"n/a"')
    assert "DCF!$C$34/(1+$A5)^1" in celula
    assert "DCF!$C$38*(1+B$4)/($A5-B$4)/(1+$A5)^5" in celula
    assert celula.endswith("-DCF!$B$8)/DCF!$B$9)")


def test_planilha_dcf_recusa_financeira_e_falta_de_dado():
    from finlab.backend import xlsx_dcf

    with pytest.raises(xlsx_dcf.SemDados, match="financeira"):
        xlsx_dcf.planilha("ITUB4", _prem_planilha(), {"financial": True})

    with pytest.raises(xlsx_dcf.SemDados, match="fluxo de caixa"):
        xlsx_dcf.planilha("XPTO3", _prem_planilha(fcf_media3=None, fcf_ultimo=None),
                          {"financial": False})

    with pytest.raises(xlsx_dcf.SemDados, match="ações"):
        xlsx_dcf.planilha("XPTO3", _prem_planilha(shares=None, shares_emitidas=None),
                          {"financial": False})

    with pytest.raises(xlsx_dcf.SemDados, match="cotação"):
        xlsx_dcf.planilha("XPTO3", _prem_planilha(preco=None),
                          {"financial": False})

    # só a média disponível: o último exercício cai para ela, e vice-versa
    blob = xlsx_dcf.planilha("XPTO3", _prem_planilha(fcf_ultimo=None),
                             {"financial": False})
    ws = _abrir_planilha(blob)["DCF"]
    assert ws["B7"].value == ws["B6"].value == 4800


# ---------------------------------------------------------------------------
# Análise de call pelo agente de contexto (spec 4.2/4.4 do redesenho)
# ---------------------------------------------------------------------------

def test_rotulo_da_call_prioriza_titulo_e_cai_no_calendario():
    """O rótulo XTXX identifica a call no arquivo e nas citações. O título
    manda; sem ele, vale o calendário de resultados: a call divulga o
    trimestre ANTERIOR à sua data — confere com o arquivo de referência."""
    from finlab.backend import call_analise as ca

    assert ca.rotulo_da_call("2026-08-07", "Call do 3T26") == "3T26"
    # heurística: agosto→2T, maio→1T, fevereiro→4T do ano anterior…
    assert ca.rotulo_da_call("2026-08-07") == "2T26"
    assert ca.rotulo_da_call("2026-05-08") == "1T26"
    assert ca.rotulo_da_call("2026-02-27") == "4T25"
    assert ca.rotulo_da_call("2025-10-30") == "3T25"
    assert ca.rotulo_da_call("2025-07-31") == "2T25"


def test_analise_de_call_valida_em_codigo_nao_no_prompt(monkeypatch):
    """O modelo pode devolver nota inventada e âncora que não existe — quem
    barra é o validador, não a instrução. Só âncoras da transcrição passam."""
    from finlab.backend import call_analise as ca, calls as bcalls

    seg = bcalls.segmentar(_CALL)
    resposta = {
        "nota": "ÓTIMA",                        # fora do conjunto → None
        "entregue": "margem de 34%, recorde",
        "devendo": None,
        "preocupacao": "prazo da desalavancagem",
        "motivo": "a call não menciona promessas anteriores",
        "trechos": ["2T26#qa-01", "2T26#qa-99", "outra-call#qa-01", 123],
    }
    monkeypatch.setattr(ca.agents, "chat",
                        lambda *a, **k: "```json\n" + json.dumps(resposta) + "\n```")

    analise = ca.analisar({"provider": "openrouter", "api_key": "k",
                           "model": "m"}, "2T26", seg)
    assert analise["nota"] is None
    assert analise["entregue"] == "margem de 34%, recorde"
    assert analise["devendo"] is None
    assert analise["motivo"] == "a call não menciona promessas anteriores"
    # só a âncora que EXISTE sobreviveu
    assert analise["trechos"] == ["2T26#qa-01"]

    # resposta sem JSON nenhum → ValueError (o chamador guarda o erro)
    monkeypatch.setattr(ca.agents, "chat", lambda *a, **k: "não sei responder")
    with pytest.raises(ValueError):
        ca.analisar({"provider": "openrouter", "api_key": "k", "model": "m"},
                    "2T26", seg)


def test_call_com_slot_analisa_cacheia_e_regenera_o_md(tmp_path, monkeypatch):
    """O fluxo inteiro da rota: indexa, UMA chamada ao provedor, cache em
    data/calls/, e o {ticker}-calls.md regenerado com o formato do arquivo
    de referência — e indexado para a mesa citar [call:XTXX]."""
    from finlab.backend import app as bapp, call_analise as ca, docs as bdocs

    monkeypatch.setattr(bdocs, "DB_PATH", tmp_path / "docs.sqlite")
    monkeypatch.setattr(ca, "DIR_CALLS", tmp_path / "calls")

    chamadas = []

    def fake_chat(provider, api_key, model, system, user, **kw):
        chamadas.append({"system": system, "user": user})
        return json.dumps({
            "nota": "positiva",
            "entregue": "margem de 34%, recorde da companhia",
            "devendo": "capex de 2027 sem detalhamento",
            "preocupacao": "prazo para desalavancar até 2,0x; resposta com data (4T27)",
            "motivo": None,
            "trechos": ["2T26#qa-01"],
        })
    monkeypatch.setattr(ca.agents, "chat", fake_chat)

    r = bapp.api_call_nova("WEGE3", {
        "data": "2026-08-07", "texto": _CALL, "titulo": "Call do 2T26",
        "slot": {"provider": "openrouter", "api_key": "k", "model": "m"},
    })
    assert r["rotulo"] == "2T26"
    assert r["analise"]["nota"] == "positiva"
    assert len(chamadas) == 1
    # o critério fixo da nota viaja no prompt — é ele que torna calls comparáveis
    assert "ENTREGOU o que havia prometido" in chamadas[0]["system"]
    assert "2T26#qa-01" in chamadas[0]["user"]

    # cache em disco: navegar entre calls nunca chama o LLM de novo
    reg = ca.carregar("WEGE3", "call-2026-08-07")
    assert reg["analise"]["nota"] == "positiva"
    lista = bapp.api_calls("WEGE3")["calls"]
    assert lista[0]["nota"] == "positiva" and lista[0]["na_tela"] is True
    assert len(chamadas) == 1, "listar chamou o LLM"

    # o arquivo canônico, no formato de referência
    md = ca.caminho_md("WEGE3").read_text(encoding="utf-8")
    assert "# WEGE3 · Arquivo de calls" in md
    assert "## [call:2T26] — 07/08/2026 · nota: **POSITIVA** · na tela" in md
    assert "**Entregue:** margem de 34%" in md
    assert "`2T26#qa-01`" in md

    # e a mesa recupera a ANÁLISE pelo índice, com o marcador [call:XTXX]
    achados = bdocs.search(universe.get("WEGE3").cd_cvm, "capex 2027 sem detalhamento")
    assert any("[call:2T26]" in a["trecho"] for a in achados)


def test_call_sem_provedor_degrada_e_arquiva_na_quarta(tmp_path, monkeypatch):
    """Sem slot, a call entra no índice sem análise — com a mensagem certa.
    E a 4ª call empurra a mais antiga para fora da tela, com a data de
    arquivamento = a data da call que a empurrou (regra do arquivo)."""
    from finlab.backend import app as bapp, call_analise as ca, docs as bdocs

    monkeypatch.setattr(bdocs, "DB_PATH", tmp_path / "docs.sqlite")
    monkeypatch.setattr(ca, "DIR_CALLS", tmp_path / "calls")

    datas = ["2025-10-30", "2026-02-27", "2026-05-08", "2026-08-07"]
    for d in datas[:3]:
        r = bapp.api_call_nova("WEGE3", {"data": d, "texto": _CALL})
        assert r["analise"] is None
        assert "provedor" in r["mensagem"].lower()

    lista = bapp.api_calls("WEGE3")["calls"]
    assert [c["na_tela"] for c in lista] == [True, True, True]

    bapp.api_call_nova("WEGE3", {"data": datas[3], "texto": _CALL})
    lista = bapp.api_calls("WEGE3")["calls"]
    assert [c["rotulo"] for c in lista] == ["2T26", "1T26", "4T25", "3T25"]
    assert [c["na_tela"] for c in lista] == [True, True, True, False]
    # 3T25 saiu quando a 2T26 entrou: arquivada na data DELA
    assert lista[3]["arquivada_em"] == "2026-08-07"

    md = ca.caminho_md("WEGE3").read_text(encoding="utf-8")
    assert "## [call:3T25] — 30/10/2025 · nota: **SEM ANÁLISE** · " \
           "arquivada da tela em 07/08/2026" in md
    assert md.count("## [call:") == 4, "o arquivo tem de guardar TODAS as calls"

    # remover por engano: o cache sai e o arquivo é regenerado sem a call
    bapp.api_call_remover("WEGE3", "call-2026-08-07")
    md = ca.caminho_md("WEGE3").read_text(encoding="utf-8")
    assert md.count("## [call:") == 3
    assert ca.carregar("WEGE3", "call-2026-08-07") is None


def test_erro_do_provedor_nao_perde_a_call(tmp_path, monkeypatch):
    """Provedor fora do ar no meio do upload: a transcrição fica indexada e a
    mensagem explica — nunca um 500 que joga o trabalho fora."""
    from finlab.backend import agents as bagents, app as bapp
    from finlab.backend import call_analise as ca, docs as bdocs

    monkeypatch.setattr(bdocs, "DB_PATH", tmp_path / "docs.sqlite")
    monkeypatch.setattr(ca, "DIR_CALLS", tmp_path / "calls")

    def explode(*a, **k):
        raise bagents.LLMError("Chave de API rejeitada (401).")
    monkeypatch.setattr(ca.agents, "chat", explode)

    r = bapp.api_call_nova("WEGE3", {
        "data": "2026-08-07", "texto": _CALL,
        "slot": {"provider": "openrouter", "api_key": "ruim", "model": "m"},
    })
    assert r["analise"] is None
    assert "401" in r["mensagem"]
    assert [c["protocolo"] for c in bapp.api_calls("WEGE3")["calls"]] \
        == ["call-2026-08-07"]


# ---------------------------------------------------------------------------
# Resposta JSON — NaN de dado faltando não pode derrubar a página
# ---------------------------------------------------------------------------

def test_resposta_troca_nan_e_inf_por_null(caplog):
    from finlab.backend.app import JSONSeguro, app
    corpo = JSONSeguro({"a": float("nan"), "b": [1.5, float("inf")],
                        "c": {"d": float("-inf"), "e": "ok"}}).body
    assert json.loads(corpo) == {"a": None, "b": [1.5, None], "c": {"d": None, "e": "ok"}}
    assert "c.d" in caplog.text and "b[1]" in caplog.text
    assert app.router.default_response_class is JSONSeguro
